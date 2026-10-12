"""Workflow spreadsheet imports through generated process definitions."""
import io
import json
from pathlib import Path
import re
import unittest

from openpyxl import load_workbook

from features.helix.helix_excel import build_excel_template, excel_spec, import_excel
from features.helix.service_generator import create_package, create_service
from features.helix.workflows import HEADERS, REFERENCE, apply_workflow, fill_workflow_sample


class WorkflowTests(unittest.TestCase):
    def setUp(self):
        self.source = json.loads((Path(__file__).parents[1] / 'reference/Workflow.json').read_text())
        fulfillment = json.loads((Path(__file__).parents[1] / 'reference/WorkOrder.json').read_text())
        for key in ('connectorIdToConnection', 'rxConnectionIdToConnectionId'):
            self.source[key] = fulfillment[key]
        self.workbook = load_workbook(io.BytesIO(build_excel_template(self.source)))
        spec = excel_spec(self.source)
        template = spec['templates'][0]
        self.service = {'templateId': template['id'], 'name': 'Workflow service', 'fieldValues': {}, 'categoryIds': []}
        values = {'Service name': self.service['name'], 'Service template': template['name']}
        for header, (template_id, field_id) in spec['columns'].items():
            if template_id == template['id']:
                values[header] = 'Profile'
                self.service['fieldValues'][field_id] = 'Profile'
        self.write('Services', 2, values)

    def tearDown(self):
        self.workbook.close()

    def write(self, sheet, number, values):
        for column, cell in enumerate(self.workbook[sheet][1], 1):
            if cell.value in values:
                self.workbook[sheet].cell(number, column, values[cell.value])

    def row(self, number, **values):
        if 'Approval type' in values:
            values['Step type'] = values.pop('Approval type')
        if 'Approval summary' in values:
            values['Summary'] = values.pop('Approval summary')
        self.write('Workflow', number, {'Service name': self.service['name'], **values})

    def imported(self):
        contents = io.BytesIO()
        self.workbook.save(contents)
        return import_excel(self.source, contents.getvalue())

    def workflow(self, service=None):
        export = create_package(self.source, [service or self.service])
        main_id = export['services'][0]['allRevisions'][0]['workflowId']
        return export, export['processDefinitionGuidToProcessDefinition'][main_id]

    def test_template_hides_automatic_nodes_and_variables(self):
        self.assertEqual([c.value for c in self.workbook['Workflow'][1][:len(HEADERS)]], HEADERS)
        self.assertIn('Label', HEADERS)
        self.assertIn('Step type', HEADERS)
        for hidden in ('From', 'To', 'Condition', 'Scope', 'Variable name'):
            self.assertNotIn(hidden, HEADERS)
        self.assertTrue(all(c.value is None for c in self.workbook['Workflow'][2]))
        result = self.imported()
        self.assertEqual(result['errors'], [])
        export, process = self.workflow(result['services'][0])
        context = process['inputParams'][0]
        self.assertEqual(context['name'], 'Context')
        self.assertEqual(context['resourceType'], REFERENCE['fields']['Context']['resourceType'])
        self.assert_pipeline(export, process, [])
        del self.workbook['Workflow']
        result = self.imported()
        self.assertEqual(result['errors'], [])
        export, process = self.workflow(result['services'][0])
        self.assert_pipeline(export, process, [])

    def assert_pipeline(self, export, process, expected):
        elements = process['flowElements']
        nodes = {e['guid']: e for e in elements if not e['resourceType'].endswith('SequenceFlowDefinition')}
        flows = [e for e in elements if e['resourceType'].endswith('SequenceFlowDefinition')]
        starts = [e for e in nodes.values() if e['resourceType'].endswith('StartEventDefinition')]
        ends = [e for e in nodes.values() if e['resourceType'].endswith('EndEventDefinition')]
        self.assertEqual(len(starts), 1)
        self.assertEqual(len(ends), len(expected) + 1 if not any(e['name'] == 'Decision' for e in nodes.values()) else 1)
        final_end = next(e for e in ends if e['name'] == 'End')
        rejection_ends = set()
        def outgoing(identifier):
            return [f for f in flows if f['sourceNode'] == identifier]
        qa = nodes[outgoing(starts[0]['guid'])[0]['targetNode']]
        self.assertEqual(qa['name'], 'QAnswer')
        self.assertEqual(len([n for n in nodes.values() if n['name'] == 'QAnswer']), 1)
        called_id = export['processDefinitionNameToProcessDefinitionGuid'][qa['calledProcessDefinitionName']]
        self.assertIn(called_id, export['processDefinitionGuidToProcessDefinition'])
        current = outgoing(qa['guid'])[0]['targetNode']
        for prefix in expected:
            approval = nodes[current]
            self.assertTrue(approval['name'].startswith(prefix))
            gateway = nodes[outgoing(current)[0]['targetNode']]
            self.assertTrue(gateway['resourceType'].endswith('ExclusiveGatewayDefinition'))
            branches = outgoing(gateway['guid'])
            self.assertEqual(len(branches), 2)
            passed = next(f for f in branches if f['condition'])
            self.assertEqual(passed['condition'], '${activityResults.' + approval['guid'] + '.Status} = "Approved"')
            rejected = nodes[next(f for f in branches if not f['condition'])['targetNode']]
            self.assertEqual(rejected['actionTypeName'], 'myit-sb:setServiceRequestStatus')
            self.assertIn({'assignTarget': 'status', 'expression': '"Rejected"'}, rejected['inputMap'])
            rejection_end = outgoing(rejected['guid'])[0]['targetNode']
            self.assertIn(rejection_end, {e['guid'] for e in ends})
            if len(ends) > 1:
                self.assertNotEqual(rejection_end, final_end['guid'])
                self.assertNotIn(rejection_end, rejection_ends)
                rejection_ends.add(rejection_end)
            current = passed['targetNode']
        while current != final_end['guid']:
            node = nodes[current]
            self.assertIn(node.get('actionTypeName'), ['myit-sb:buildInputSet', 'com.bmc.myservice.connect.remedy:createRemedyWorkOrderWithIdentifiers'])
            self.assertEqual(len(outgoing(current)), 1)
            current = outgoing(current)[0]['targetNode']
        self.assertEqual(current, final_end['guid'])
        self.assertEqual({c['id'] for c in json.loads(process['layout'])['cells']}, {e['guid'][3:] for e in elements})

    def test_ordered_approvals_generate_status_checks_and_independent_dependencies(self):
        fill_workflow_sample(self.workbook, self.service['name'])
        result = self.imported()
        self.assertEqual(result['errors'], [])
        export, process = self.workflow(result['services'][0])
        self.assert_pipeline(export, process, ['LineManagerApproval', 'CustomApproval', 'GroupApproval'])
        adhoc = next(e for e in process['flowElements'] if e['name'].startswith('CustomApproval') and e.get('inputMap'))
        self.assertIn({'assignTarget': 'Flex Char 01', 'expression': '"Max"'}, adhoc['inputMap'])
        services = [result['services'][0], {**result['services'][0], 'name': 'Second service'}]
        package = create_package(self.source, services)
        self.assertEqual(len(package['processDefinitionGuidToProcessDefinition']), 4)
        self.assertEqual(len(package['processDefinitionNameToProcessDefinitionGuid']), 4)
        for definition in package['processDefinitionGuidToProcessDefinition'].values():
            for scope in ('inputParams', 'outputParams', 'localVariables'):
                self.assertTrue(all(re.fullmatch('[A-Z][A-Za-z]*', p['name']) for p in definition[scope]))

    def test_approval_order_and_summary_are_user_controlled(self):
        self.row(2, **{'Approval type': 'Custom Approval', 'Approvers': 'Custom user', 'Approval summary': 'Custom "summary"'})
        self.row(3, **{'Approval type': 'Line manager approval'})
        result = self.imported()
        self.assertEqual(result['errors'], [])
        export, process = self.workflow(result['services'][0])
        self.assert_pipeline(export, process, ['CustomApproval', 'LineManagerApproval'])
        approval = next(e for e in process['flowElements'] if e['name'].startswith('CustomApproval') and e.get('inputMap'))
        self.assertIn({'assignTarget': 'Approval Summary', 'expression': json.dumps('Custom "summary"')}, approval['inputMap'])

    def test_invalid_approval_rows_prevent_import(self):
        for values in [{'Approval type': 'Start'}, {'Approval type': 'QAnswer'},
                       {'Approval type': 'Custom Approval'}, {'Approval type': 'Line manager approval', 'Approvers': 'User'}]:
            with self.subTest(values=values):
                for c in self.workbook['Workflow'][2]:
                    c.value = None
                self.row(2, **values)
                result = self.imported()
                self.assertEqual(result['services'], [])
                self.assertTrue(any(e['sheet'] == 'Workflow' and e['row'] == 2 for e in result['errors']))

    def test_older_workflows_move_qanswer_immediately_after_start(self):
        qa_target = next(iter(REFERENCE['dependencies'].values()))['name']
        rows = [
            {'Row type': 'Step', 'Name': 'Start', 'Type': 'Start'},
            {'Row type': 'Step', 'Name': 'LineManagerApproval', 'Type': 'Call process', 'Target': 'com.bmc.arsys.myitsb.myitsb-approval:Request Approval'},
            {'Row type': 'Step', 'Name': 'Decision', 'Type': 'Exclusive gateway'},
            {'Row type': 'Step', 'Name': 'QAnswers', 'Type': 'Call process', 'Target': qa_target, 'Input mapping': 'SBContext = ${processContext.Context}'},
            {'Row type': 'Step', 'Name': 'Reject', 'Type': 'Service task', 'Target': 'myit-sb:setServiceRequestStatus', 'Input mapping': 'status = "Rejected"'},
            {'Row type': 'Step', 'Name': 'End', 'Type': 'End'},
            {'Row type': 'Flow', 'From': 'Start', 'To': 'LineManagerApproval'},
            {'Row type': 'Flow', 'From': 'LineManagerApproval', 'To': 'Decision'},
            {'Row type': 'Flow', 'From': 'Decision', 'To': 'QAnswers', 'Condition': '${activityResults.LineManagerApproval.Status} = "Approved"'},
            {'Row type': 'Flow', 'From': 'Decision', 'To': 'Reject'},
            {'Row type': 'Flow', 'From': 'Reject', 'To': 'End'},
            {'Row type': 'Flow', 'From': 'QAnswers', 'To': 'End'},
        ]
        export, process = self.workflow({**self.service, 'workflow': rows})
        self.assert_pipeline(export, process, ['LineManagerApproval'])

    def test_unreachable_steps_and_invalid_gateway_branches_are_rejected(self):
        steps = [{'Row type': 'Step', 'Name': name, 'Type': kind} for name, kind in
                 [('Start', 'Start'), ('Decision', 'Exclusive gateway'), ('Approved', 'End'), ('Rejected', 'End')]]
        flows = [{'Row type': 'Flow', 'From': 'Start', 'To': 'Decision'},
                 {'Row type': 'Flow', 'From': 'Decision', 'To': 'Approved', 'Condition': 'true'},
                 {'Row type': 'Flow', 'From': 'Decision', 'To': 'Rejected'}]
        def process():
            return {'flowElements': [], 'layout': '{"cells": []}', 'inputParams': [], 'outputParams': [], 'localVariables': [],
                    'lastUpdateTime': 123, 'lastChangedBy': None, 'owner': None,
                    'tags': None, 'description': None, 'internal': False}
        apply_workflow(process(), steps + flows, 123)
        with self.assertRaisesRegex(ValueError, 'reachable'):
            apply_workflow(process(), steps + flows[:-1], 123)
        with self.assertRaisesRegex(ValueError, 'default Condition'):
            apply_workflow(process(), steps + [dict(row, Condition='true') for row in flows if row['From'] == 'Decision'] + flows[:1], 123)
        bad = [dict(row, Condition='${processContext.Missing} = "Yes"') if row.get('Condition') else row for row in flows]
        with self.assertRaisesRegex(ValueError, 'unknown variable'):
            apply_workflow(process(), steps + bad, 123)

    def test_context_cannot_be_redefined_and_question_names_are_pascal(self):
        export = create_service(self.source, self.service['templateId'], 'Service', self.service['fieldValues'], [], questions=[
            {'type': 'TextField', 'label': 'request summary!'}, {'type': 'TextField', 'label': 'request-summary'},
            {'type': 'TextField', 'label': '123'}])
        process = export['processDefinitionGuidToProcessDefinition'][export['services'][0]['allRevisions'][0]['workflowId']]
        self.assertEqual([p['name'] for p in process['inputParams']], ['RequestSummary', 'RequestSummaryNext', 'OneTwoThree', 'Context'])
        for row in [{'Row type': 'Variable', 'Name': 'Context', 'Type': 'Text', 'Scope': 'Local'},
                    {'Row type': 'Variable', 'Name': 'Other', 'Type': 'Context', 'Scope': 'Input'}]:
            with self.assertRaisesRegex(ValueError, 'reserved'):
                apply_workflow({'inputParams': [], 'outputParams': [], 'localVariables': []}, [row], 123)


if __name__ == '__main__':
    unittest.main()
