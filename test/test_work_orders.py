"""Reference-based fulfillment steps, mappings, and exported connection metadata."""
import io
import json
from pathlib import Path
import unittest

from openpyxl import load_workbook

from features.helix.helix_excel import build_excel_template, excel_spec, import_excel
from features.helix.service_generator import create_package
from features.helix.workflows import HEADERS, PREVIOUS_HEADERS


class WorkOrderTests(unittest.TestCase):
    def setUp(self):
        self.source = json.loads((Path(__file__).parents[1] / 'reference/WorkOrder.json').read_text())
        self.workbook = load_workbook(io.BytesIO(build_excel_template(self.source)))
        spec = excel_spec(self.source)
        template = spec['templates'][0]
        values = {'Service name': 'Service', 'Service template': template['name']}
        for header, (template_id, _) in spec['columns'].items():
            if template_id == template['id']:
                values[header] = 'Profile'
        self.write('Services', 2, values)

    def tearDown(self):
        self.workbook.close()

    def write(self, sheet, number, values):
        for column, cell in enumerate(self.workbook[sheet][1], 1):
            if cell.value in values:
                self.workbook[sheet].cell(number, column, values[cell.value])

    def row(self, number, **values):
        self.write('Workflow', number, {'Service name': 'Service', **values})

    def imported(self):
        contents = io.BytesIO()
        self.workbook.save(contents)
        return import_excel(self.source, contents.getvalue())

    def generated(self):
        result = self.imported()
        self.assertEqual(result['errors'], [])
        export = create_package(self.source, result['services'])
        process = export['processDefinitionGuidToProcessDefinition'][export['services'][0]['allRevisions'][0]['workflowId']]
        return export, process

    @staticmethod
    def inputs(element):
        return {item['assignTarget']: item['expression'] for item in element['inputMap']}

    def test_group_approval_and_unique_labels(self):
        self.write('Questions', 2, {'Service name': 'Service', 'Label': 'Approver group', 'Type': 'Short text'})
        self.row(2, **{'Label': 'Support team approval', 'Step type': 'Group Approval', 'Approvers': 'question: Approver group'})
        self.row(3, **{'Step type': 'Group Approval', 'Approvers': 'Example group'})
        _, process = self.generated()
        names = [e['name'].casefold() for e in process['flowElements']]
        self.assertEqual(len(names), len(set(names)))
        group = next(e for e in process['flowElements'] if e['name'] == 'Support team approval')
        self.assertEqual(self.inputs(group)['Flex Char 01'], '(${processContext.ApproverGroup} + " | Request Approver")')
        checks = [e for e in process['flowElements'] if e.get('condition')]
        self.assertTrue(any('${activityResults.' + group['guid'] + '.Status}' in e['condition'] for e in checks))
        self.row(3, **{'Label': 'SUPPORT TEAM APPROVAL'})
        result = self.imported()
        self.assertEqual(result['services'], [])
        self.assertTrue(any(e['row'] == 3 and 'unique' in e['message'] for e in result['errors']))

    def test_work_order_with_identifiers_and_dynamic_input_set(self):
        self.write('Questions', 2, {'Service name': 'Service', 'Label': 'Priority', 'Type': 'Dropdown', 'Choices': 'Medium\nHigh'})
        self.row(2, **{'Label': 'Team approval', 'Step type': 'Group Approval', 'Approvers': 'Example group'})
        self.row(3, **{'Label': 'Create team work order', 'Step type': 'Create work order',
                      'Support company': 'Example company', 'Support organization': 'Example organization',
                      'Support group': 'Example group', 'Priority': 'question: Priority', 'Other fields': 'SLA_Code = 11X6'})
        export, process = self.generated()
        items = process['flowElements']
        wo = next(e for e in items if e.get('actionTypeName', '').endswith(':createRemedyWorkOrderWithIdentifiers'))
        builder = next(e for e in items if e.get('actionTypeName') == 'myit-sb:buildInputSet')
        layout = {c['id']: c for c in json.loads(process['layout'])['cells']}
        self.assertEqual(layout[wo['guid'][3:]]['type'], 'rx.ProcessAction')
        self.assertEqual(layout[builder['guid'][3:]]['type'], 'rx.ProcessAction')
        qa = next(e for e in items if e['name'] == 'QAnswer')
        mappings = self.inputs(wo)
        self.assertEqual(mappings['description'], '${activityResults.' + qa['guid'] + '.OutputText}')
        self.assertEqual(mappings['otherFields'], '${activityResults.' + builder['guid'] + '.output.inputs}')
        self.assertEqual(mappings['requestAssigneeSupportGroupCompany'], '"Example company"')
        self.assertEqual(mappings['requestAssigneeSupportOrganization'], '"Example organization"')
        self.assertEqual(mappings['requestAssigneeSupportGroupName'], '"Example group"')
        self.assertNotIn('templateName', mappings)
        self.assertEqual(self.inputs(builder), {'key1': '"Priority"', 'value1': '${processContext.Priority}', 'key2': '"SLA_Code"', 'value2': '"11X6"'})
        self.assertTrue(any(e.get('sourceNode') == builder['guid'] and e.get('targetNode') == wo['guid'] for e in items))
        # Approved approval flow targets the builder, and never bypasses it.
        self.assertTrue(any(e.get('targetNode') == builder['guid'] and 'Approved' in e.get('condition', '') for e in items))
        status = next(p for p in process['localVariables'] if p['name'] == 'WorkStatus')
        self.assertTrue(status['resourceType'].endswith('CharacterFieldDefinition'))
        self.assertFalse(any(p['name'] == 'WorkStatus' for p in process['inputParams']))
        identifier = json.loads(mappings['connectionId'])
        self.assertEqual(export['rxConnectionIdToConnectionId'][identifier], self.source['rxConnectionIdToConnectionId'][identifier])
        key = str(export['rxConnectionIdToConnectionId'][identifier])
        self.assertEqual(export['connectorIdToConnection'][key], self.source['connectorIdToConnection'][key])
        called_id = export['processDefinitionNameToProcessDefinitionGuid'][qa['calledProcessDefinitionName']]
        called = export['processDefinitionGuidToProcessDefinition'][called_id]
        self.assertTrue(any(p['name'] == 'OutputText' for p in called['outputParams']))
        self.assertEqual(len({e['name'].casefold() for e in items}), len(items))

    def test_template_only_and_no_extra_fields_skip_builder(self):
        self.row(2, **{'Step type': 'Create work order', 'Template': 'User-supplied template'})
        _, process = self.generated()
        wo = next(e for e in process['flowElements'] if e.get('actionTypeName', '').endswith(':createRemedyWorkOrderWithIdentifiers'))
        self.assertEqual(self.inputs(wo)['templateName'], '"User-supplied template"')
        self.assertNotIn('otherFields', self.inputs(wo))
        self.assertFalse(any(e.get('actionTypeName') == 'myit-sb:buildInputSet' for e in process['flowElements']))
        self.assertFalse(any('requestAssignee' in k for k in self.inputs(wo)))

    def test_invalid_work_order_fields_prevent_entire_import(self):
        cases = [
            {}, {'Support company': 'Company'},
            {'Template': 'Template', 'Support group': 'Group'},
            {'Template': 'Template', 'Connection ID': 'unknown'},
            {'Template': 'Template', 'Priority': 'Medium', 'Other fields': 'Priority = High'},
            {'Template': 'Template', 'Other fields': 'description = override'},
            {'Template': 'Template', 'Priority': 'question: Missing'},
        ]
        for fields in cases:
            with self.subTest(fields=fields):
                for cell in self.workbook['Workflow'][2]:
                    cell.value = None
                self.row(2, **{'Step type': 'Create work order', **fields})
                result = self.imported()
                self.assertEqual(result['services'], [])
                self.assertTrue(any(e['sheet'] == 'Workflow' for e in result['errors']))

    def test_local_work_status_for_legacy_incident_and_change_creation(self):
        result = self.imported()
        service = result['services'][0]
        for target in ('example:createIncident', 'example:createChange', 'example:createWO'):
            rows = [
                {'Row type': 'Step', 'Name': 'Start', 'Type': 'Start'},
                {'Row type': 'Step', 'Name': 'CreateItem', 'Type': 'Service task', 'Target': target},
                {'Row type': 'Step', 'Name': 'End', 'Type': 'End'},
                {'Row type': 'Flow', 'From': 'Start', 'To': 'CreateItem'},
                {'Row type': 'Flow', 'From': 'CreateItem', 'To': 'End'},
            ]
            export = create_package(self.source, [{**service, 'workflow': rows}])
            process = export['processDefinitionGuidToProcessDefinition'][export['services'][0]['allRevisions'][0]['workflowId']]
            self.assertEqual([p['name'] for p in process['localVariables']], ['WorkStatus'])

    def test_question_name_collision_and_connector_merging(self):
        self.write('Questions', 2, {'Service name': 'Service', 'Label': 'Work Status', 'Type': 'Short text'})
        self.row(2, **{'Step type': 'Create work order', 'Template': 'Template', 'Priority': 'question: Work Status'})
        result = self.imported()
        self.assertEqual(result['errors'], [])
        service = result['services'][0]
        export = create_package(self.source, [{**service, 'name': 'Approval only', 'workflow': []}, service])
        main = export['processDefinitionGuidToProcessDefinition'][export['services'][1]['allRevisions'][0]['workflowId']]
        builder = next(e for e in main['flowElements'] if e.get('actionTypeName') == 'myit-sb:buildInputSet')
        self.assertEqual(self.inputs(builder)['value1'], '${processContext.WorkStatusNext}')
        self.assertEqual([p['name'] for p in main['localVariables']], ['WorkStatus'])
        self.assertTrue(any(p['name'] == 'WorkStatusNext' for p in main['inputParams']))
        self.assertEqual(len(export['rxConnectionIdToConnectionId']), 1)
        self.assertEqual(len(export['connectorIdToConnection']), 1)

    def test_hidden_default_connection_and_service_dropdowns(self):
        from features.helix.fulfillment import resolve_connection
        self.assertTrue(self.workbook['Workflow'].column_dimensions['J'].hidden)
        defined = self.workbook.defined_names['ServiceNameOptions'].attr_text
        self.assertIn("'Services'!$A$2:$A$1000", defined)
        self.assertIn('LOOKUP', defined)
        for name in ('Questions', 'Workflow'):
            validation = next(v for v in self.workbook[name].data_validations.dataValidation if str(v.sqref) == 'A2:A1001')
            self.assertEqual(validation.formula1, 'ServiceNameOptions')
            self.assertFalse(validation.showDropDown)
        environment = {'connectorIdToConnection': {
            'a': {'rxBundleId': 'com.bmc.myservice.connect.remedy', 'connectionName': 'Secondary'},
            'b': {'rxBundleId': 'com.bmc.myservice.connect.remedy', 'connectionName': 'Remedy'}},
            'rxConnectionIdToConnectionId': {'secondary-id': 'a', 'default-id': 'b'}}
        self.assertEqual(resolve_connection({}, environment), 'default-id')

    def test_sample_always_includes_work_order_even_without_existing_workflows(self):
        from features.helix.workflows import fill_workflow_sample
        from features.helix.fulfillment import REFERENCE
        environment = {**self.source, 'processDefinitionGuidToProcessDefinition': {}}
        fill_workflow_sample(self.workbook, 'Service', environment)
        rows = [dict(zip(HEADERS, [c.value or '' for c in row[:len(HEADERS)]]))
                for row in self.workbook['Workflow'].iter_rows(min_row=2, max_row=5)]
        self.assertEqual([r['Step type'] for r in rows], ['Line manager approval', 'Custom Approval', 'Group Approval', 'Create work order'])
        reference = {m['assignTarget']: m['expression'] for m in REFERENCE['workOrder']['inputMap']}
        self.assertEqual(rows[-1]['Template'], json.loads(reference['templateName']))
        self.assertEqual(rows[-1]['Priority'], 'Medium')
        self.assertEqual(rows[-1]['Other fields'], 'SLA_Code = 11X6\nASCPY = DIB-IT')
        self.assertEqual(rows[-1]['Connection ID'], '')
        self.assertEqual(self.imported()['errors'], [])
        # A sample stays complete when the environment has no Remedy metadata;
        # import then reports the missing connection instead of dropping fulfillment.
        self.source['rxConnectionIdToConnectionId'] = {}
        fill_workflow_sample(self.workbook, 'Service', {})
        self.assertEqual(self.workbook['Workflow'].cell(5, HEADERS.index('Step type') + 1).value, 'Create work order')
        self.assertTrue(any('Remedy' in e['message'] for e in self.imported()['errors']))

    def test_previous_approval_only_template_is_accepted(self):
        del self.workbook['Workflow']
        sheet = self.workbook.create_sheet('Workflow')
        sheet.append(PREVIOUS_HEADERS)
        sheet.append(['Service', 'Custom Approval', 'Summary', 'Example approver'])
        self.assertEqual(self.imported()['errors'], [])


if __name__ == '__main__':
    unittest.main()
