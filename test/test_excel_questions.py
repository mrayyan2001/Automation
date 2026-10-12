"""Exercise XLSX question input through import and generated Helix exports."""
import io
import json
from pathlib import Path
import unittest
from datetime import datetime

from openpyxl import load_workbook

from features.helix.helix_excel import build_excel_template, import_excel, excel_spec
from features.helix.excel_questions import QUESTION_HEADERS, CHOICE_HEADERS, TABLE_HEADERS, LEGACY_SHEETS, SIMPLE_TYPES
from features.helix.helix_service import HelixService
from features.helix.excel_questions import QUESTION_TYPES
import tempfile
from features.helix.service_generator import create_package


class ExcelQuestionsTest(unittest.TestCase):
    def setUp(self):
        self.source = json.loads((Path(__file__).parents[1] / 'reference/Service with questions.json').read_text())
        self.workbook = load_workbook(io.BytesIO(build_excel_template(self.source)))
        # Exercise compatibility with the previous multi-sheet format.
        del self.workbook['Questions']
        for name, headers in LEGACY_SHEETS.items():
            self.workbook.create_sheet(name).append(headers)
        spec = excel_spec(self.source)
        template = spec['templates'][0]
        self.service_row = {'Service name': 'Excel service', 'Service template': template['name']}
        for header, (template_id, field_id) in spec['columns'].items():
            if template_id == template['id']:
                field = next(f for f in template['fields'] if f['id'] == field_id)
                self.service_row[header] = field.get('params', {}).get('form_values', [{}])[0].get('name', 'Profile') if field['type'] == 'DROPDOWN' else 'Profile'
        self.write('Services', 2, self.service_row)

    def tearDown(self):
        self.workbook.close()

    def write(self, name, row, values):
        headers = [c.value for c in self.workbook[name][1]]
        for column, title in enumerate(headers, 1):
            if title in values:
                self.workbook[name].cell(row, column, values[title])

    def question(self, row, key, kind, **fields):
        self.write('Questions', row, {'Service name': 'Excel service', 'Question key': key, 'Type': kind, 'Label': key, **fields})

    def contents(self):
        output = io.BytesIO()
        self.workbook.save(output)
        return output.getvalue()

    def test_questions_choices_tables_groups_defaults_and_conditions_export(self):
        self.question(2, 'name', 'Short text', **{'Required': 'Yes', 'Default answer': 'Alice', 'Maximum characters': 100})
        self.question(3, 'department', 'Dropdown', **{'Show when question key': 'name', 'Expected answers': 'Alice'})
        self.question(4, 'employees', 'Multiple selection table', **{'Table columns': 'Name\nEmail'})
        self.question(5, 'items', 'Repeatable group', **{'Group limit': 4})
        self.question(6, 'item count', 'Number', **{'Parent group key': 'items', 'Default answer': 0, 'Minimum number': 0})
        for row, label, value, default in [(2, 'HR', 'hr', 'Yes'), (3, 'IT', 'it', 'No')]:
            self.write('Choices', row, {'Service name': 'Excel service', 'Question key': 'department', 'Label': label, 'Saved value': value, 'Default': default})
        self.write('Table rows', 2, {'Service name': 'Excel service', 'Question key': 'employees', 'Saved row value': '001', 'Default': 'Yes', 'Value 1': 'Alice', 'Value 2': 'alice@example.com'})
        result = import_excel(self.source, self.contents())
        self.assertEqual(result['errors'], [])
        service = result['services'][0]
        self.assertEqual(len(service['questions']), 4)
        department = service['questions'][1]
        self.assertEqual(department['config']['defaultAnswer'], 'hr')
        self.assertEqual(department['config']['visibilityCondition']['leadingQuestionId'], service['questions'][0]['id'])
        table = service['questions'][2]['config']
        self.assertEqual(table['questionMultiColumn']['data'][0]['displayValue'], ['Alice', 'alice@example.com'])
        self.assertEqual(table['defaultAnswers'], ['001'])
        self.assertEqual(service['questions'][3]['pageItems'][0]['config']['defaultAnswer'], 0)
        export = create_package(self.source, result['services'])
        representation = json.loads(export['questionnaires'][0]['jsonRepresentation'])
        params = {p['id']: p['name'] for p in export['processDefinitionGuidToProcessDefinition'][representation['workflowId']]['inputParams']}
        for q in representation['pages'][0]['pageItems']:
            self.assertEqual(params[q['processInputId']], __import__('features.helix.workflows', fromlist=['pascal_name']).pascal_name(q['label']))
        self.assertEqual(export['connectorIdToConnection'], {})

    def test_old_services_only_workbook_is_supported(self):
        for name in ('Questions', 'Choices', 'Table rows', 'Conditions', 'Guide'):
            del self.workbook[name]
        result = import_excel(self.source, self.contents())
        self.assertEqual(result['errors'], [])
        self.assertNotIn('questions', result['services'][0])

    def test_bad_question_or_orphaned_detail_prevents_whole_import(self):
        for sheet, values in [
            ('Questions', {'Service name': 'Unknown', 'Question key': 'q', 'Type': 'Short text', 'Label': 'Q'}),
            ('Choices', {'Service name': 'Excel service', 'Question key': 'missing', 'Label': 'A'}),
            ('Table rows', {'Service name': 'Excel service', 'Question key': 'missing', 'Saved row value': '1'}),
        ]:
            with self.subTest(sheet=sheet):
                self.write(sheet, 2, values)
                result = import_excel(self.source, self.contents())
                self.assertEqual(result['services'], [])
                self.assertTrue(any(e['sheet'] == sheet and e['row'] == 2 for e in result['errors']))
                for cell in self.workbook[sheet][2]:
                    cell.value = None

    def test_choices_and_table_cells_reject_duplicates_and_formulas(self):
        self.question(2, 'choice', 'Dropdown')
        for row in (2, 3):
            self.write('Choices', row, {'Service name': 'Excel service', 'Question key': 'choice', 'Label': 'Same', 'Saved value': 'same'})
        result = import_excel(self.source, self.contents())
        self.assertEqual(result['services'], [])
        self.assertTrue(any(e['sheet'] == 'Choices' and e['row'] == 3 for e in result['errors']))
        self.workbook['Choices']['C3'] = '=1+1'
        result = import_excel(self.source, self.contents())
        self.assertTrue(any(e['sheet'] == 'Choices' and 'formulas' in e['message'] for e in result['errors']))

    def test_group_cycles_and_missing_children_are_rejected(self):
        self.question(2, 'group', 'Repeatable group', **{'Parent group key': 'group'})
        result = import_excel(self.source, self.contents())
        self.assertEqual(result['services'], [])
        self.assertTrue(any('nested groups' in e['message'] for e in result['errors']))

    def test_same_key_can_be_used_in_separate_services(self):
        self.write('Services', 3, {**self.service_row, 'Service name': 'Second'})
        self.question(2, 'name', 'Short text')
        self.write('Questions', 3, {'Service name': 'Second', 'Question key': 'name', 'Type': 'TextField', 'Label': 'name'})
        result = import_excel(self.source, self.contents())
        self.assertEqual(result['errors'], [])
        self.assertEqual([len(s['questions']) for s in result['services']], [1, 1])
        self.assertNotEqual(result['services'][0]['questions'][0]['id'], result['services'][1]['questions'][0]['id'])

    def test_typed_excel_dates_and_relative_limits(self):
        self.question(2, 'start date', 'Date', **{'Default answer': datetime(2026, 10, 12),
                      'Earliest date': datetime(2026, 10, 9), 'Latest date': 'today+30'})
        result = import_excel(self.source, self.contents())
        self.assertEqual(result['errors'], [])
        config = result['services'][0]['questions'][0]['config']
        self.assertEqual(config['defaultAnswer'], '2026-10-12')
        self.assertEqual(config['min']['specificLimit'], '2026-10-09')
        self.assertEqual(config['max']['relativeLimit']['numberOfUnits'], 30)

    def test_simple_template_and_sample_round_trip(self):
        with tempfile.TemporaryDirectory() as directory:
            storage = HelixService(Path(directory) / 'helix.sqlite3')
            environment = storage.save_environment('Sample', self.source)
            sample = storage.sample_service(environment['id'])
            contents = build_excel_template(self.source, sample)
        workbook = load_workbook(io.BytesIO(contents))
        self.assertEqual(workbook.sheetnames, ['Services', 'Questions', 'Workflow', 'Guide'])
        self.assertTrue(workbook['Services'].column_dimensions['B'].hidden)
        self.assertEqual([c.value for c in workbook['Questions'][1][:len(QUESTION_HEADERS)]], QUESTION_HEADERS)
        operation_list = next(v for v in workbook['Questions'].data_validations.dataValidation if str(v.sqref) == 'L2:L1001')
        self.assertLess(len(operation_list.formula1), 256)
        self.assertEqual(operation_list.formula1, 'ConditionOperations')
        source = workbook.defined_names['ConditionOperations']
        source_sheet, source_range = next(source.destinations)
        values = [cell.value for row in workbook[source_sheet][source_range] for cell in row]
        self.assertEqual(values, ['is', 'is not', 'contains', 'does not contain', 'starts with', 'ends with', 'includes', 'does not include'])
        self.assertEqual(workbook['Questions']['DF2'].value, '(no value)')
        value_list = next(v for v in workbook['Questions'].data_validations.dataValidation if str(v.sqref) == 'M2')
        self.assertFalse(value_list.showErrorMessage)
        self.assertFalse(value_list.showDropDown)
        self.assertIn('COUNTIF', value_list.formula1)
        self.assertIn('$DF$2:$HB$2', value_list.formula1)
        workbook.close()
        result = import_excel(self.source, contents)
        self.assertEqual(result['errors'], [])
        questions = result['services'][0]['questions']
        self.assertEqual({q['type'] for q in questions}, set(SIMPLE_TYPES.values()))
        for q in questions:
            for option in q['config'].get('options', []):
                self.assertEqual(option['dataValue'], option['displayValue'])
                self.assertFalse(option['disabled'])
        self.assertEqual(create_package(self.source, result['services'])['connectorIdToConnection'], {})

    def test_inline_choices_conditions_and_row_errors(self):
        del self.workbook['Questions']
        for name in ('Choices', 'Table rows', 'Conditions'):
            del self.workbook[name]
        self.workbook.create_sheet('Questions').append(QUESTION_HEADERS)
        self.write('Questions', 2, {'Service name': 'Excel service', 'Label': 'Department', 'Type': 'Dropdown', 'Choices': 'HR\nIT'})
        self.write('Questions', 3, {'Service name': 'Excel service', 'Label': 'Details', 'Type': 'Short text', 'Show if label': 'Department', 'Operation': 'is', 'Value': 'HR'})
        result = import_excel(self.source, self.contents())
        self.assertEqual(result['errors'], [])
        questions = result['services'][0]['questions']
        self.assertEqual(questions[1]['config']['visibilityCondition']['leadingQuestionId'], questions[0]['id'])
        self.write('Questions', 2, {'Choices': 'HR\nHR'})
        result = import_excel(self.source, self.contents())
        self.assertEqual(result['services'], [])
        self.assertTrue(any(e['sheet'] == 'Questions' and e['row'] == 2 for e in result['errors']))

    def test_condition_operations_and_value_kinds_export(self):
        from features.helix.excel_questions import TEXT_OPERATIONS, MULTI_OPERATIONS
        del self.workbook['Questions']
        for name in ('Choices', 'Table rows', 'Conditions'):
            del self.workbook[name]
        sheet = self.workbook.create_sheet('Questions')
        sheet.append(QUESTION_HEADERS)
        leaders = [('Text', 'Short text', ''), ('Single', 'Dropdown', 'One\nTwo'),
                   ('Multi', 'Multiple choice dropdown', 'One\nTwo')]
        for number, (label, kind, choices) in enumerate(leaders, 2):
            self.write('Questions', number, {'Service name': 'Excel service', 'Label': label, 'Type': kind, 'Choices': choices})
        number = 5
        expected = {}
        for leader, _, _ in leaders:
            operations = MULTI_OPERATIONS if leader == 'Multi' else TEXT_OPERATIONS
            for operation, operator in operations.items():
                for value in ('(no value)', 'custom text', 'One'):
                    label = f'Dependent {number}'
                    self.write('Questions', number, {'Service name': 'Excel service', 'Label': label, 'Type': 'Short text',
                               'Show if label': leader, 'Operation': operation, 'Value': value})
                    expected[label] = (leader, operator, value)
                    number += 1
        result = import_excel(self.source, self.contents())
        self.assertEqual(result['errors'], [])
        export = create_package(self.source, result['services'])
        items = json.loads(export['questionnaires'][0]['jsonRepresentation'])['pages'][0]['pageItems']
        lookup = {q['label']: q for q in items}
        for label, (leader, operator, value) in expected.items():
            condition = lookup[label]['visibilityCondition']
            expression = condition['expression']
            self.assertEqual(expression['operator'], operator)
            self.assertEqual(expression['left']['valueReference'], lookup[leader]['id'])
            right = expression['right']
            selected = value == 'One' and leader != 'Text'
            self.assertEqual(right['type'], 'QUESTION_OPTION' if selected else 'STRING')
            self.assertEqual(right['label'], 'One' if selected else None)
            self.assertEqual(right['value'], None if value == '(no value)' else value)
            self.assertEqual(right['values'], [] if value == '(no value)' else [value])
        # Type restrictions use the leading question, not the dependent question.
        self.write('Questions', 5, {'Show if label': 'Multi', 'Operation': 'contains'})
        result = import_excel(self.source, self.contents())
        self.assertEqual(result['services'], [])
        self.assertTrue(any(e['row'] == 5 and 'includes' in e['message'] for e in result['errors']))

    def test_reference_condition_null_and_choice_shapes(self):
        reference = json.loads((Path(__file__).parents[1] / 'reference/Conditions.json').read_text())
        items = json.loads(reference['questionnaires'][0]['jsonRepresentation'])['pages'][0]['pageItems']
        null_right = items[1]['visibilityCondition']['expression']['left']['right']
        self.assertEqual(null_right, {'values': [], 'type': 'STRING', 'label': None,
                                    'valueReference': None, 'fieldType': None, 'value': None})
        choice_right = items[3]['visibilityCondition']['expression']['right']
        self.assertEqual(choice_right['type'], 'QUESTION_OPTION')
        self.assertEqual(choice_right['values'], ['2'])

    def test_invalid_condition_tree_prevents_import(self):
        self.question(2, 'leader', 'Short text')
        self.question(3, 'dependent', 'Short text')
        self.write('Conditions', 2, {'Service name': 'Excel service', 'Question key': 'dependent',
                   'Rule key': 'root', 'Match': 'All'})
        self.write('Conditions', 3, {'Service name': 'Excel service', 'Question key': 'dependent',
                   'Rule key': 'child', 'Parent rule key': 'root', 'Leading question key': 'leader',
                   'Comparison': 'Equals', 'Expected answers': 'test'})
        result = import_excel(self.source, self.contents())
        self.assertEqual(result['services'], [])
        self.assertTrue(any(e['sheet'] == 'Conditions' and e['row'] == 2 for e in result['errors']))
