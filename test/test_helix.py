"""Exercise the environment workflow through the application's controller layer."""

import io
import asyncio
import json
import tempfile
import unittest
import zipfile
from pathlib import Path
from unittest.mock import patch
from uuid import uuid4

import httpx
from fastapi import FastAPI
from starlette.exceptions import HTTPException

from common.exception.global_exception_handler import http_exception_handler
from context import ApplicationContext
from features.helix.helix_router import HelixRouter
from features.helix.helix_service import HelixService
from features.helix.service_generator import category_hierarchy, create_service
from features.helix.helix_excel import excel_spec
from openpyxl import load_workbook

class CategoryHierarchyTest(unittest.TestCase):
    def test_parent_links_group_categories_independent_of_order(self):
        categories = {
            "3": {"name": "Laptop", "parentId": 2},
            "5": {"name": "Accounts", "parentId": "1"},
            "2": {"name": "Hardware", "parentId": "1"},
            "1": {"name": "IT", "parentId": None},
            "4": {"name": "Monitor", "parentId": "2"},
            "6": {"name": "Other", "parentId": "missing"},
        }
        hierarchy = category_hierarchy(categories)
        self.assertEqual([c["id"] for c in hierarchy], ["1", "5", "2", "3", "4", "6"])
        laptop = next(c for c in hierarchy if c["id"] == "3")
        self.assertEqual(laptop["path"], "IT / Hardware / Laptop")
        self.assertEqual(laptop["parentPath"], "IT / Hardware")
        self.assertEqual(len(category_hierarchy({"1": {"name": "A", "parentId": "2"}, "2": {"name": "B", "parentId": "1"}})), 2)


ROOT = Path(__file__).resolve().parents[1]


class HelixTests(unittest.IsolatedAsyncioTestCase):
    async def test_download_excel_sample_and_import(self):
        source = json.loads((ROOT / 'reference' / 'Service with questions.json').read_text())
        fulfillment = json.loads((ROOT / 'reference' / 'WorkOrder.json').read_text())
        for key in ('connectorIdToConnection', 'rxConnectionIdToConnectionId'):
            source[key] = fulfillment[key]
        environment = self.service.save_environment('Excel sample', source)
        base = f"/helix/api/environments/{environment['id']}"
        response = await self.client.get(base + '/excel-sample')
        self.assertEqual(response.status_code, 200, response.text)
        self.assertIn('Helix-services-sample.xlsx', response.headers['content-disposition'])
        imported = await self.client.post(base + '/import-excel', content=response.content)
        self.assertEqual(imported.status_code, 200, imported.text)
        data = imported.json()['data']
        self.assertEqual(data['errors'], [])
        self.assertEqual(len(data['services'][0]['questions']), 15)
        package = await self.client.post(base + '/packages', json={'services': data['services']})
        self.assertEqual(package.status_code, 200, package.text)

    async def asyncSetUp(self):
        self.directory = tempfile.TemporaryDirectory()
        self.database = Path(self.directory.name) / "helix.sqlite3"
        self.service = HelixService(self.database)
        app = FastAPI()
        app.add_exception_handler(HTTPException, http_exception_handler)
        app.include_router(HelixRouter.__router__)
        app.state.context = ApplicationContext()
        app.state.context.beans[HelixService] = self.service
        self.client = httpx.AsyncClient(transport=httpx.ASGITransport(app=app), base_url="http://test")

    async def asyncTearDown(self):
        await self.client.aclose()
        self.directory.cleanup()

    async def test_empty_exports_create_and_download_service_packages(self):
        for index in (1, 3):
            with self.subTest(export=index):
                source = json.loads((ROOT / "reference" / f"ServiceExport {index}.json").read_text())
                source["services"] = []
                saved = await self.client.post("/helix/api/environments", json={"name": f"Environment {index}", "export": source})
                self.assertEqual(saved.status_code, 200, saved.text)
                environment = saved.json()["data"]
                identifier = environment["id"]
                template_id = environment["templates"][0]["id"]
                generated = await self.client.post(f"/helix/api/environments/{identifier}/services", json={"templateId": template_id, "name": "Laptop Request", "fieldValues": {f["id"]: "Requested profile text" for f in environment["templates"][0]["required"]}, "categoryIds": []})
                self.assertEqual(generated.status_code, 200, generated.text)
                package = generated.json()["data"]
                url = f"/helix/api/environments/{identifier}/packages/{package['id']}"
                download = await self.client.get(url)
                export = download.json()
                self.assertEqual(len(export["services"]), 1)
                self.assertEqual(export["versionInfo"], source["versionInfo"])
                self.assertEqual(export["serviceTypes"], source["serviceTypes"])
                self.assertEqual(export["categories"], source["categories"])
                identities = set()
                for service in export["services"]:
                    revision = service["versions"][0]["revisions"][0]
                    self.assertEqual(service["allRevisions"], [revision])
                    self.assertEqual(service["name"], "Laptop Request")
                    self.assertEqual(revision["state"], "draft")
                    self.assertEqual(revision["prices"][0]["paymentType"], "FREE")
                    self.assertNotIn("questionnaireGroupId", revision)
                    self.assertNotIn("logoId", revision)
                    self.assertIn(service["id"], export["selectedServiceIds"])
                    self.assertIn(revision["workflowId"], export["processDefinitionGuidToProcessDefinition"])
                    self.assertEqual({f["serviceTypeFieldId"] for f in revision["fieldValues"]},
                                     {f["id"] for f in environment["templates"][0]["required"]})
                    for obj in (service, service["versions"][0], revision):
                        self.assertNotIn(obj["guid"], identities)
                        identities.add(obj["guid"])
                self.assertFalse(export["oldRevisionGuidToLogo"])
                self.assertFalse(export["questionnaires"])
                zip_response = await self.client.get(url + "?format=zip")
                with zipfile.ZipFile(io.BytesIO(zip_response.content)) as archive:
                    self.assertEqual(json.loads(archive.read("ServiceExport.json")), export)
                # A fresh storage instance reads the saved environment and package after restart.
                restored = HelixService(self.database).get_environment(identifier)
                self.assertEqual(restored["packages"][0]["id"], package["id"])
                second = await self.client.post(f"/helix/api/environments/{identifier}/services", json={"templateId": template_id, "name": "VPN Request", "fieldValues": {f["id"]: "Second profile text" for f in environment["templates"][0]["required"]}})
                self.assertTrue(set(second.json()["data"]["names"]).isdisjoint(package["names"]))
                missing = await self.client.get(f"/helix/api/environments/{uuid4()}/packages/{package['id']}")
                self.assertEqual(missing.status_code, 404)
        self.assertEqual(len((await self.client.get("/helix/api/environments")).json()["data"]), 2)

    async def test_invalid_upload_duplicate_and_template_errors(self):
        for body in ({"name": "Bad", "export": {}}, {"name": " ", "export": {}}, {"name": "Bad", "export": {"serviceTypes": []}}):
            response = await self.client.post("/helix/api/environments", json=body)
            self.assertEqual(response.status_code, 400)
        response = await self.client.post("/helix/api/environments", content=b"x" * (5 * 1024 * 1024 + 1))
        self.assertEqual(response.status_code, 413)
        source = json.loads((ROOT / "reference" / "ServiceExport 1.json").read_text())
        environment = self.service.save_environment("Development", source)
        duplicate = await self.client.post("/helix/api/environments", json={"name": " development ", "export": source})
        self.assertEqual(duplicate.status_code, 409)
        response = await self.client.post(f"/helix/api/environments/{environment['id']}/services", json={"templateId": "missing", "name": "Test"})
        self.assertEqual(response.status_code, 400)
        self.assertEqual(self.service.get_environment(environment["id"])["packages"], [])
        self.assertEqual((await self.client.get("/helix/")).status_code, 200)

    async def test_parent_category_can_be_selected_and_exported(self):
        source = json.loads((ROOT / "reference" / "ServiceExport 3.json").read_text())
        source["categories"] = {
            "parent": {"name": "IT", "parentId": None},
            "child": {"name": "Hardware", "parentId": "parent"},
            "other": {"name": "Unrelated", "parentId": None},
        }
        for template in source["serviceTypes"].values():
            template["categories"] = ["child"]
        environment = self.service.save_environment("Parent categories", source)
        template = environment["templates"][0]
        self.assertEqual(template["categoryIds"], ["parent", "child"])
        spec = excel_spec(source)
        self.assertIn(["parent", "IT"], spec["categories"])
        values = {field["id"]: "Profile text" for field in template["required"]}
        response = await self.client.post(f"/helix/api/environments/{environment['id']}/services", json={
            "name": "Parent service", "templateId": template["id"], "fieldValues": values, "categoryIds": ["parent"]})
        self.assertEqual(response.status_code, 200, response.text)
        export = self.service.get_package(environment["id"], response.json()["data"]["id"])
        self.assertEqual(export["services"][0]["allRevisions"][0]["categoryIds"], ["parent"])
        self.assertEqual(export["serviceTypes"], source["serviceTypes"])
        with self.assertRaises(ValueError):
            create_service(source, template["id"], "Invalid", values, ["other"])

    async def test_form_fields_categories_and_validation(self):
        source = json.loads((ROOT / "reference" / "ServiceExport 3.json").read_text())
        environment = self.service.save_environment("IT Request", source)
        identifier = environment["id"]
        detail = (await self.client.get(f"/helix/api/environments/{identifier}")).json()["data"]
        template = detail["templates"][0]
        self.assertIn("fields", template)
        self.assertTrue(detail["categories"])
        endpoint = f"/helix/api/environments/{identifier}/services"
        values = {f["id"]: "User supplied text" for f in template["required"]}
        category = next(c["id"] for c in detail["categories"] if c["id"] in template["categoryIds"])
        body = {"name": "New IT Request", "templateId": template["id"], "fieldValues": values, "categoryIds": [category]}
        for change in ({"fieldValues": {}}, {"name": " "}, {"categoryIds": ["missing"]},
                       {"fieldValues": {**values, "missing": "x"}}, {"tags": [" "]}, {"tags": ["x" * 101]}):
            response = await self.client.post(endpoint, json={**body, **change})
            self.assertEqual(response.status_code, 400, response.text)
        self.assertEqual(self.service.get_environment(identifier)["packages"], [])
        rich = next(f for f in template["fields"] if f["type"] == "RICH_INPUT")
        body["fieldValues"] = {**values, rich["id"]: "<script>alert(1)</script>\nSecond line"}
        body.pop("templateId")
        body["tags"] = ["t1", " T1 ", "t2", "New tag"]
        response = await self.client.post(endpoint, json=body)
        self.assertEqual(response.status_code, 200, response.text)
        package = self.service.get_package(identifier, response.json()["data"]["id"])
        revision = package["services"][0]["allRevisions"][0]
        fields = {f["serviceTypeFieldId"]: f["value"][0] for f in revision["fieldValues"]}
        self.assertEqual(fields[rich["id"]], "<p>&lt;script&gt;alert(1)&lt;/script&gt;<br>Second line</p>")
        self.assertEqual(revision["categoryIds"], [category])
        self.assertEqual(package["services"][0]["serviceTypeId"], template["id"])
        self.assertEqual(package["tags"]["1601"], source["tags"]["1601"])
        self.assertEqual(package["tags"]["1602"], source["tags"]["1602"])
        self.assertEqual(len(package["tags"]), 3)
        self.assertEqual({t["name"] for t in package["tags"].values()}, {"t1", "t2", "New tag"})
        self.assertEqual({t["tagId"] for t in revision["revisionTags"]}, set(package["tags"]))
        self.assertTrue(all(t["summary"] == "SbeServiceRevisionTag" and t["locale"] is None for t in revision["revisionTags"]))
        self.assertEqual(revision, package["services"][0]["versions"][0]["revisions"][0])

    async def test_multi_service_package_and_atomic_validation(self):
        source = json.loads((ROOT / "reference" / "ServiceExport 3.json").read_text())
        desktop = json.loads((ROOT / "reference" / "ServiceExport 2.json").read_text())
        source["serviceTypes"].update(desktop["serviceTypes"])
        environment = self.service.save_environment("Mixed package", source)
        endpoint = f"/helix/api/environments/{environment['id']}/packages"
        services = [
            {"name": "Laptop", "templateId": "1", "fieldValues": {"1": "Laptop excerpt", "2": "Laptop description"}, "tags": ["t1", "shared new tag"]},
            {"name": "Desktop", "templateId": "2", "fieldValues": {"3": "Desktop excerpt", "5": "Desktop description"}, "tags": ["shared new tag", "t2"]},
        ]
        invalid = await self.client.post(endpoint, json={"services": [services[0], {**services[1], "fieldValues": {}}]})
        self.assertEqual(invalid.status_code, 400)
        duplicate = await self.client.post(endpoint, json={"services": [services[0], {**services[1], "name": " laptop "}]})
        self.assertEqual(duplicate.status_code, 400)
        self.assertEqual(self.service.get_environment(environment["id"])["packages"], [])
        response = await self.client.post(endpoint, json={"services": services})
        self.assertEqual(response.status_code, 200, response.text)
        summary = response.json()["data"]
        self.assertEqual(summary["names"], ["Laptop", "Desktop"])
        url = f"/helix/api/environments/{environment['id']}/packages/{summary['id']}"
        export = (await self.client.get(url)).json()
        self.assertEqual(len(export["services"]), 2)
        self.assertEqual(len(set(export["selectedServiceIds"])), 2)
        self.assertEqual(len(export["processDefinitionGuidToProcessDefinition"]), 2)
        self.assertEqual(len(export["tags"]), 3)
        shared_id = next(k for k, t in export["tags"].items() if t["name"] == "shared new tag")
        for service, template_id in zip(export["services"], ["1", "2"]):
            self.assertEqual(service["serviceTypeId"], template_id)
            revision = service["allRevisions"][0]
            self.assertEqual(revision, service["versions"][0]["revisions"][0])
            self.assertIn(shared_id, {t["tagId"] for t in revision["revisionTags"]})
            self.assertTrue(all(t["tagId"] in export["tags"] for t in revision["revisionTags"]))
            self.assertIn(revision["workflowId"], export["processDefinitionGuidToProcessDefinition"])
            self.assertIn(service["guid"], export["serviceTranslations"])
        with zipfile.ZipFile(io.BytesIO((await self.client.get(url + "?format=zip")).content)) as archive:
            self.assertEqual(json.loads(archive.read("ServiceExport.json")), export)
        self.assertEqual(len(HelixService(self.database).get_environment(environment["id"])["packages"]), 1)

    async def test_excel_template_import_and_package(self):
        source = json.loads((ROOT / "reference" / "ServiceExport 3.json").read_text())
        environment = self.service.save_environment("Excel environment", source)
        base = f"/helix/api/environments/{environment['id']}"
        # IDE/server processes may not have Node or any command on PATH.
        with patch.dict("os.environ", {"PATH": "", "HELIX_SPREADSHEET_MODULES": "/missing/runtime"}):
            download = await self.client.get(base + "/excel-template")
        self.assertEqual(download.status_code, 200, download.text if download.status_code != 200 else "")
        workbook = load_workbook(io.BytesIO(download.content))
        self.assertEqual(workbook.sheetnames, ["Services", "Questions", "Choices", "Table rows", "Conditions", "Guide"])
        self.assertEqual(workbook["Services"].freeze_panes, "A2")
        self.assertIn("ServicesTable", workbook["Services"].tables)
        self.assertTrue(workbook["Services"].data_validations.dataValidation)
        spec = excel_spec(source)
        self.assertEqual([workbook["Services"].cell(1, c).value for c in range(1, len(spec["headers"]) + 1)], spec["headers"])
        validations = workbook["Services"].data_validations.dataValidation
        self.assertEqual({v.formula1 for v in validations}, {"ServiceTemplateOptions", "ServiceCategoryOptions"})
        self.assertTrue(all(v.showDropDown is False for v in validations))
        category_validation = next(v for v in validations if v.formula1 == "ServiceCategoryOptions")
        self.assertEqual(str(category_validation.sqref), "C2:E51")
        self.assertTrue(all(workbook["Services"].cell(row, column).fill.patternType is None for row in (1, 2) for column in range(1, len(spec["headers"]) + 1)))
        self.assertTrue(workbook["Services"].column_dimensions["CX"].hidden)
        self.assertTrue(workbook["Services"].column_dimensions["CY"].hidden)
        workbook.close()

        def file_with_rows(rows):
            fixture = load_workbook(io.BytesIO(download.content))
            for row, values in enumerate(rows, 2):
                for column, value in enumerate(values, 1):
                    fixture["Services"].cell(row, column, value)
            output = io.BytesIO()
            fixture.save(output)
            fixture.close()
            return output.getvalue()

        valid = ["Laptop", "IT Request", "723", "326", "1", "t1,shared", "Laptop excerpt", "Laptop description"]
        second = ["Monitor", "1", "1", "", "", "t2,shared", "Monitor excerpt", "Monitor description"]
        contents = await asyncio.to_thread(file_with_rows, [valid, second])
        imported = await self.client.post(base + "/import-excel", content=contents)
        self.assertEqual(imported.status_code, 200, imported.text)
        result = imported.json()["data"]
        self.assertFalse(result["errors"])
        self.assertEqual(len(result["services"]), 2)
        self.assertEqual(result["services"][0]["categoryIds"], ["723", "326", "1"])
        self.assertEqual(self.service.get_environment(environment["id"])["packages"], [])
        saved = await self.client.post(base + "/packages", json={"services": result["services"]})
        self.assertEqual(saved.status_code, 200, saved.text)
        export = self.service.get_package(environment["id"], saved.json()["data"]["id"])
        self.assertEqual([s["name"] for s in export["services"]], ["Laptop", "Monitor"])

        too_many = ["Too many categories", "IT Request", "723;326", "1", "4", "", "Excerpt", "Description"]
        overflow = await self.client.post(base + "/import-excel", content=await asyncio.to_thread(file_with_rows, [too_many]))
        self.assertEqual(overflow.status_code, 200)
        self.assertFalse(overflow.json()["data"]["services"])
        self.assertIn("3 categories", overflow.json()["data"]["errors"][0]["message"])
        invalid = ["Broken", "IT Request", "", "", "", "", "", "Description"]
        response = await self.client.post(base + "/import-excel", content=await asyncio.to_thread(file_with_rows, [valid, invalid]))
        result = response.json()["data"]
        self.assertFalse(result["services"])
        self.assertEqual(result["errors"][0]["row"], 3)
        self.assertIn("required", result["errors"][0]["message"])
        self.assertEqual(len(self.service.get_environment(environment["id"])["packages"]), 1)
        bad_file = await self.client.post(base + "/import-excel", content=b"not an xlsx")
        self.assertEqual(bad_file.status_code, 400)

    async def test_environment_templates_and_no_platform(self):
        source = json.loads((ROOT / "reference" / "ServiceExport 2.json").read_text())
        environment = self.service.save_environment("Desktop only", source)
        endpoint = f"/helix/api/environments/{environment['id']}/services"
        body = {"name": "Desktop request", "templateId": "2", "fieldValues": {"3": "Excerpt", "5": "Description"}}
        response = await self.client.post(endpoint, json=body)
        self.assertEqual(response.status_code, 200, response.text)
        package = self.service.get_package(environment["id"], response.json()["data"]["id"])
        self.assertEqual(package["services"][0]["serviceTypeId"], "2")
        rejected = await self.client.post(endpoint, json={**body, "fieldValues": {**body["fieldValues"], "6": "Windows"}})
        self.assertEqual(rejected.status_code, 400)
        it_source = json.loads((ROOT / "reference" / "ServiceExport 1.json").read_text())
        source["serviceTypes"].update(it_source["serviceTypes"])
        source["categories"].update(it_source["categories"])
        combined = self.service.save_environment("Combined", source)
        endpoint = f"/helix/api/environments/{combined['id']}/services"
        for template_id, values in (("1", {"1": "Excerpt", "2": "Description"}), ("2", {"3": "Excerpt", "5": "Description"})):
            response = await self.client.post(endpoint, json={"name": "Service " + template_id, "templateId": template_id, "fieldValues": values})
            self.assertEqual(response.status_code, 200, response.text)
            package = self.service.get_package(combined["id"], response.json()["data"]["id"])
            self.assertEqual(package["services"][0]["serviceTypeId"], template_id)
        missing = await self.client.post(endpoint, json={"name": "Missing template", "fieldValues": {}})
        self.assertEqual(missing.status_code, 400)


if __name__ == "__main__":
    unittest.main()
