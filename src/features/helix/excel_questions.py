"""Plain-cell XLSX question definitions, choices, and static table rows."""
from datetime import date, datetime
from html import escape
from uuid import uuid4

from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.worksheet.table import Table, TableStyleInfo

QUESTION_TYPES = {
    'Short text': 'TextField', 'Long text': 'TextArea', 'Number': 'IntegerField',
    'Date': 'DateField', 'Date and time': 'DateTimeField', 'Yes / No': 'YesNo',
    'Dropdown': 'Dropdown', 'Radio buttons': 'RadioButtons', 'Checkboxes': 'Checkbox',
    'Multiple choice dropdown': 'MultiSelectDropdown', 'Table': 'DataTable',
    'Multiple selection table': 'MultiSelectDataTable', 'Attachments': 'Attachment',
    'Repeatable group': 'RepeatableGroup', 'Description': 'DescriptionWidget',
    'Section heading': 'SectionWidget',
}
CHOICE_TYPES = {'YesNo', 'Dropdown', 'RadioButtons', 'Checkbox', 'MultiSelectDropdown'}
TABLE_TYPES = {'DataTable', 'MultiSelectDataTable'}
MULTI_TYPES = {'Checkbox', 'MultiSelectDropdown', 'MultiSelectDataTable', 'Attachment'}
LEGACY_QUESTION_HEADERS = [
    'Service name', 'Question key', 'Parent group key', 'Type', 'Label', 'Required',
    'Help text', 'Default answer', 'Hidden', 'Read-only', 'Maximum characters',
    'Minimum number', 'Maximum number', 'Number step', 'Text rows', 'Rich text',
    'Table columns', 'Group limit', 'Group output', 'Timezone', 'Earliest date',
    'Latest date', 'Validation pattern', 'Validation message', 'Show when question key',
    'Comparison', 'Expected answers',
]
CHOICE_HEADERS = ['Service name', 'Question key', 'Label', 'Saved value', 'Default', 'Disabled']
TABLE_HEADERS = ['Service name', 'Question key', 'Saved row value', 'Default'] + [f'Value {n}' for n in range(1, 21)]
CONDITION_HEADERS = ['Service name', 'Question key', 'Rule key', 'Parent rule key', 'Match',
                     'Leading question key', 'Comparison', 'Expected answers', 'Group parentheses']
LEGACY_SHEETS = {'Questions': LEGACY_QUESTION_HEADERS, 'Choices': CHOICE_HEADERS, 'Table rows': TABLE_HEADERS,
          'Conditions': CONDITION_HEADERS}

QUESTION_HEADERS = ['Service name', 'Label', 'Type', 'Required', 'Default answer', 'Hidden', 'Read-only', 'Earliest date', 'Latest date', 'Choices', 'Show if label', 'Operation', 'Value']
SHEETS = {'Questions': QUESTION_HEADERS}
SIMPLE_TYPES = {label: kind for label, kind in QUESTION_TYPES.items() if kind not in TABLE_TYPES | {'RepeatableGroup'}}

TEXT_OPERATIONS = {'is': 'EQUALS', 'is not': 'NOT_EQUALS', 'contains': 'CONTAINS',
                   'does not contain': 'NOT_CONTAINS', 'starts with': 'STARTS_WITH', 'ends with': 'ENDS_WITH'}
MULTI_OPERATIONS = {'includes': 'INCLUDES', 'does not include': 'NOT_INCLUDES'}
PREVIOUS_QUESTION_HEADERS = QUESTION_HEADERS[:11] + ['Answer is']


def add_question_sheets(workbook):
    for name, headers in SHEETS.items():
        sheet = workbook.create_sheet(name)
        sheet.sheet_view.showGridLines = False
        sheet.append(headers)
        sheet.freeze_panes = 'C2'
        sheet.row_dimensions[1].height = 42
        sheet.sheet_format.defaultRowHeight = 36 if name == 'Questions' else 24
        for column, title in enumerate(headers, 1):
            sheet.column_dimensions[get_column_letter(column)].width = 28 if column <= 5 else 24
            for row in range(1, 102):
                cell = sheet.cell(row, column)
                cell.number_format = '@'
                cell.font = Font(name='Arial', size=11, bold=row == 1, color='000000')
                cell.alignment = Alignment(vertical='center', horizontal='center' if row == 1 else 'left', wrap_text=True)
        table = Table(displayName=name.replace(' ', '') + 'Table', ref=f'A1:{get_column_letter(len(headers))}101')
        table.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
        sheet.add_table(table)
        sheet.row_dimensions[2].height = 36 if name == 'Questions' else 24

    def dropdown(sheet_name, column, values, source_name):
        sheet = workbook[sheet_name]
        # Sources are text, outside the editable table and importer range.
        source_column = 102
        while sheet.cell(1, source_column).value is not None:
            source_column += 1
        letter = get_column_letter(source_column)
        for row, value in enumerate(values, 1):
            cell = sheet.cell(row, source_column, value)
            cell.data_type = 's'
        sheet.column_dimensions[letter].hidden = True
        from openpyxl.workbook.defined_name import DefinedName
        workbook.defined_names.add(DefinedName(source_name, attr_text=f"'{sheet_name}'!${letter}$1:${letter}${len(values)}"))
        validation = DataValidation(type='list', formula1=source_name, allow_blank=True)
        validation.showDropDown = False
        validation.showErrorMessage = True
        validation.error = 'Choose a value from the dropdown.'
        sheet.add_data_validation(validation)
        validation.add(f'{column}2:{column}1001')

    dropdown('Questions', 'C', list(SIMPLE_TYPES), 'QuestionTypeOptions')
    for index, column in enumerate(('D', 'F', 'G')):
        dropdown('Questions', column, ['Yes', 'No'], f'QuestionBooleanOptions{index}')
    dropdown('Questions', 'L', list(TEXT_OPERATIONS) + list(MULTI_OPERATIONS), 'ConditionOperations')
    sheet = workbook['Questions']
    leader_row = 'MATCH($A2&CHAR(1)&$K2,INDEX($A$2:$A$1001&CHAR(1)&$B$2:$B$1001,0),0)'
    # Per-row dropdown sources update when labels or choice lists are edited.
    # Free text remains allowed; import resolves known labels to QUESTION_OPTION.
    for row_number in range(2, 102):
        sheet.cell(row_number, 110, '(no value)')
        source = f'IFERROR(INDEX($J$2:$J$1001,{leader_row.replace("$A2", f"$A{row_number}").replace("$K2", f"$K{row_number}")}),"")'
        sheet.cell(row_number, 108, f'={source}')
        for option in range(1, 101):
            choice_list = f'CHAR(10)&$DD{row_number}&CHAR(10)'
            start = f'FIND(CHAR(1),SUBSTITUTE({choice_list},CHAR(10),CHAR(1),{option}))'
            end = f'FIND(CHAR(1),SUBSTITUTE({choice_list},CHAR(10),CHAR(1),{option + 1}))'
            sheet.cell(row_number, 110 + option, f'=IFERROR(TRIM(MID($DD{row_number},{start},{end}-{start}-1)),"")')
        value_validation = DataValidation(type='list', formula1=f'OFFSET($DF${row_number},0,0,1,MAX(1,COUNTIF($DF${row_number}:$HB${row_number},"?*")))', allow_blank=True)
        value_validation.showDropDown = False
        value_validation.showErrorMessage = False
        sheet.add_data_validation(value_validation)
        value_validation.add(f'M{row_number}')
    for column in range(108, 211):
        sheet.column_dimensions[get_column_letter(column)].hidden = True
    guide = workbook.create_sheet('Guide')
    guide.sheet_view.showGridLines = False
    rows = [
        ['Sheet / field', 'How to fill it'],
        ['Services', 'One service per row. Service template is hidden and defaults to the first environment template. Unhide it to select another template.'],
        ['Questions', 'One question per row. Service name must match Services. Labels must be unique within each service. Row order is question order.'],
        ['Choices', 'Enter choice labels on separate lines (Alt+Enter) in the question row. Each label is also the saved value.'],
        ['Default answer', 'Use a choice label, or one label per line for multiple selections. Numbers must be whole numbers.'],
        ['Earliest / Latest date', 'Leave blank for no limit; use YYYY-MM-DD or today+N / today-N. Date defaults use YYYY-MM-DD; date-time defaults use YYYY-MM-DDTHH:MM.'],
        ['Conditions', 'Use Show if label, Operation, and Value in the question row. The Operation dropdown lists all eight operations. Text and single choice: is, is not, contains, does not contain, starts with, ends with. Multi select: includes, does not include.'],
        ['Condition value', 'Choose a referenced choice label, type free text, or select (no value). Blank Value also means no value when Show if label is filled. Blank condition fields mean always visible. Value dropdowns show the first 100 labels for the first 100 question rows; other labels can be typed.'],
        ['Optional settings', 'Required, Hidden, and Read-only accept Yes/No. Leave unused fields blank.'],
        ['Import', 'Upload the workbook and review the package draft. Invalid rows prevent the entire import; errors identify the sheet and row.'],
        ['Limits', 'Up to 50 services, 100 questions per service, and 1,000 input rows per sheet. Keep headers unchanged.'],
    ]
    for row in rows:
        guide.append(row)
    guide.column_dimensions['A'].width = 28
    guide.column_dimensions['B'].width = 100
    for row in guide:
        for cell in row:
            cell.font = Font(name='Arial', size=11, bold=cell.row == 1)
            cell.alignment = Alignment(vertical='top', wrap_text=True)
        guide.row_dimensions[row[0].row].height = 55 if row[0].row > 1 else 28


def _rows(workbook, name, errors):
    if name not in workbook.sheetnames:
        return []
    expected = LEGACY_SHEETS[name]
    rows = workbook[name].iter_rows(max_row=1002, max_col=101)
    headers = [str(c.value).strip() if c.value is not None else '' for c in next(rows)]
    while headers and not headers[-1]:
        headers.pop()
    if headers != expected:
        raise ValueError(f'{name}: headers do not match. Download a new template and keep its headers.')
    result = []
    for number, cells in enumerate(rows, 2):
        if all(c.value is None or str(c.value).strip() == '' for c in cells):
            continue
        try:
            if number > 1001:
                raise ValueError('Use at most 1,000 input rows per sheet.')
            if any(c.value is not None for c in cells[len(expected):]):
                raise ValueError('A filled cell has no matching column header.')
            if any(c.data_type in ('f', 'e') for c in cells):
                raise ValueError('Use plain values, not formulas or Excel errors.')
            def text(value):
                if isinstance(value, (date, datetime)):
                    return value.isoformat()
                return str(value).strip() if value is not None else ''
            result.append((number, dict(zip(expected, (text(c.value) for c in cells)))))
        except ValueError as exc:
            errors.append({'sheet': name, 'row': number, 'message': str(exc)})
    return result


def _boolean(value):
    if value.casefold() in ('', 'no', 'false', '0'):
        return False
    if value.casefold() in ('yes', 'true', '1'):
        return True
    raise ValueError('Use Yes or No for boolean fields.')


def _integer(value, title, minimum=None):
    try:
        number = int(value)
    except ValueError as exc:
        raise ValueError(f'{title} must be a whole number.') from exc
    if minimum is not None and number < minimum:
        raise ValueError(f'{title} must be at least {minimum}.')
    return number


def _lines(value):
    return [v.strip() for v in value.splitlines() if v.strip()]


def _limit(value):
    if value.lower().startswith('today'):
        days = _integer(value[5:] or '0', 'Relative days')
        return {'limitType': 'RELATIVE', 'relativeLimit': {'questionId': None, 'dateUnit': 'DAYS_FROM_NOW', 'numberOfUnits': days, 'repeatableGroupInstanceId': None}}
    parsed = datetime.fromisoformat(value).date() if 'T' in value else date.fromisoformat(value)
    return {'limitType': 'SPECIFIC', 'specificLimit': parsed.isoformat()}


def _import_legacy_question_sheets(workbook, services, errors):
    """Attach validated question trees, resolving workbook keys to fresh IDs."""
    service_map = {s['name'].casefold(): s for s in services}
    questions, locations, parents, visibility = {}, {}, {}, {}
    labels = set()

    def owner(row):
        service_name = row['Service name'].casefold()
        if service_name not in service_map:
            raise ValueError('Service name must match a valid row on Services.')
        if not row['Question key']:
            raise ValueError('Enter a Question key.')
        return service_name, row['Question key']

    for number, row in _rows(workbook, 'Questions', errors):
        try:
            key = owner(row)
            if key in questions:
                raise ValueError('Question key must be unique within the service.')
            kind = QUESTION_TYPES.get(row['Type'], row['Type'])
            if kind not in QUESTION_TYPES.values():
                raise ValueError('Choose a supported static question type.')
            if not row['Label'] or len(row['Label']) > 200:
                raise ValueError('Enter a question label (1–200 characters).')
            label_key = (key[0], row['Label'].casefold())
            if kind not in ('DescriptionWidget', 'SectionWidget') and label_key in labels:
                raise ValueError('Question labels must be unique within the service, including group questions.')
            config = {}
            for column, setting in [('Hidden', 'hidden'), ('Read-only', 'readOnly'), ('Rich text', 'isRTF')]:
                if row[column]:
                    config[setting] = _boolean(row[column])
            if row['Help text']:
                config['description'] = '<p>' + escape(row['Help text']).replace('\n', '<br>') + '</p>'
            for column, setting, minimum in [('Maximum characters', 'maxLength', 1), ('Minimum number', 'minValue', None), ('Maximum number', 'maxValue', None), ('Number step', 'step', 1), ('Text rows', 'numberOfRows', 1), ('Group limit', 'instanceLimit', 1)]:
                if row[column]:
                    config[setting] = _integer(row[column], column, minimum)
            if config.get('minValue') is not None and config.get('maxValue') is not None and config['minValue'] > config['maxValue']:
                raise ValueError('Minimum number must not exceed Maximum number.')
            for column, setting in [('Timezone', 'timezoneId'), ('Validation pattern', 'regex'), ('Validation message', 'regexCustomErrorMessage')]:
                if row[column]:
                    config[setting] = row[column]
            if row['Validation message']:
                config['regexErrorMessageType'] = 'CustomMessage'
            if row['Default answer']:
                if kind == 'IntegerField':
                    config['defaultAnswer'] = _integer(row['Default answer'], 'Default answer')
                elif kind == 'DateField':
                    value = row['Default answer']
                    config['defaultAnswer'] = (datetime.fromisoformat(value).date() if 'T' in value else date.fromisoformat(value)).isoformat()
                elif kind == 'DateTimeField':
                    config['defaultAnswer'] = datetime.fromisoformat(row['Default answer']).isoformat(timespec='minutes')
                elif kind in MULTI_TYPES:
                    config['defaultAnswers'] = _lines(row['Default answer'])
                else:
                    config['defaultAnswer'] = row['Default answer']
            for column, setting in [('Earliest date', 'min'), ('Latest date', 'max')]:
                if row[column]:
                    config[setting] = _limit(row[column])
            if kind == 'RepeatableGroup':
                config['outputFormat'] = row['Group output'] or 'CSV'
                if config['outputFormat'] not in ('CSV', 'JSON'):
                    raise ValueError('Group output must be CSV or JSON.')
            if kind in TABLE_TYPES:
                columns = _lines(row['Table columns'])
                if not 1 <= len(columns) <= 20:
                    raise ValueError('Tables need 1–20 column titles, one per line in Table columns.')
                config['questionMultiColumn'] = {'metaData': [{'displayId': str(uuid4()), 'displayName': title, 'datasetFieldType': None, 'selectionItems': None, 'dataValueField': False} for title in columns], 'data': []}
            q = {'id': str(uuid4()), 'type': kind, 'label': row['Label'], 'required': _boolean(row['Required']), 'config': config}
            if kind == 'RepeatableGroup':
                q['pageItems'] = []
            questions[key] = q
            if kind not in ('DescriptionWidget', 'SectionWidget'):
                labels.add(label_key)
            locations[key] = number
            parents[key] = row['Parent group key']
            if row['Show when question key']:
                visibility[key] = (row['Show when question key'], row['Comparison'] or 'Equals', _lines(row['Expected answers']))
        except ValueError as exc:
            errors.append({'sheet': 'Questions', 'row': number, 'message': str(exc)})

    for sheet in ('Choices', 'Table rows'):
        for number, row in _rows(workbook, sheet, errors):
            try:
                key = owner(row)
                q = questions.get(key)
                if q is None:
                    raise ValueError('Question key must match a valid row on Questions.')
                c = q['config']
                if sheet == 'Choices':
                    if q['type'] not in CHOICE_TYPES:
                        raise ValueError('Choices must reference a choice question.')
                    if not row['Label']:
                        raise ValueError('Enter a choice label.')
                    value = row['Saved value'] or row['Label']
                    options = c.setdefault('options', [])
                    if any(o['dataValue'] == value for o in options):
                        raise ValueError('Saved choice values must be unique within the question.')
                    option = {'dataValue': value, 'displayValue': row['Label'], 'disabled': _boolean(row['Disabled'])}
                    options.append(option)
                    c['staticDataValues'] = [o['dataValue'] for o in options]
                else:
                    if q['type'] not in TABLE_TYPES:
                        raise ValueError('Table rows must reference a table question.')
                    table = c['questionMultiColumn']
                    value = row['Saved row value']
                    if not value or any(r['dataValue'] == value for r in table['data']):
                        raise ValueError('Each table row needs a unique Saved row value.')
                    count = len(table['metaData'])
                    if any(row[f'Value {n}'] for n in range(count + 1, 21)):
                        raise ValueError('A table cell has no matching column title on Questions.')
                    table['data'].append({'dataValue': value, 'displayValue': [row[f'Value {n}'] for n in range(1, count + 1)]})
                    c['staticDataValues'] = [r['dataValue'] for r in table['data']]
                if _boolean(row['Default']):
                    if q['type'] in MULTI_TYPES:
                        defaults = c.setdefault('defaultAnswers', [])
                        if value not in defaults:
                            defaults.append(value)
                    else:
                        if c.get('defaultAnswer') is not None and c['defaultAnswer'] != value:
                            raise ValueError('Single-selection questions can have only one default.')
                        c['defaultAnswer'] = value
            except ValueError as exc:
                errors.append({'sheet': sheet, 'row': number, 'message': str(exc)})

    comparisons = {'Equals': 'EQUALS', 'Does not equal': 'NOT_EQUALS', 'Includes': 'INCLUDES', 'Does not include': 'NOT_INCLUDES', 'Starts with': 'STARTS_WITH', 'Ends with': 'ENDS_WITH'}
    for key, q in questions.items():
        try:
            parent = questions.get((key[0], parents[key])) if parents[key] else None
            if parents[key] and (parent is None or parent['type'] != 'RepeatableGroup' or q['type'] == 'RepeatableGroup'):
                raise ValueError('Parent group key must reference a Repeatable group; nested groups are not supported.')
            if q['type'] in CHOICE_TYPES and q['type'] != 'YesNo' and not q['config'].get('options'):
                raise ValueError('Add at least one choice on Choices.')
            if q['type'] in TABLE_TYPES and not q['config']['questionMultiColumn']['data']:
                raise ValueError('Add at least one row on Table rows.')
            values = set(q['config'].get('staticDataValues', []))
            if q['type'] == 'YesNo' and not values:
                values = {'true', 'false'}
            defaults = q['config'].get('defaultAnswers', []) if q['type'] in MULTI_TYPES else [q['config']['defaultAnswer']] if q['config'].get('defaultAnswer') is not None else []
            if q['type'] in CHOICE_TYPES | TABLE_TYPES and any(v not in values for v in defaults):
                raise ValueError('Default answers must match saved choice/table values.')
            if key in visibility:
                leader_key, comparison, expected = visibility[key]
                leader = questions.get((key[0], leader_key))
                if leader is None or leader is q:
                    raise ValueError('Show when question key must reference another question in this service.')
                if comparison not in comparisons:
                    raise ValueError('Choose a supported Comparison.')
                operator = comparisons[comparison]
                q['config']['visibilityCondition'] = {'leadingQuestionId': leader['id'], 'leadingQuestionExpectedAnswers': expected, 'expression': {'operator': operator, 'withParentheses': False, 'left': {'values': [], 'type': 'QUESTION_ANSWER', 'label': leader['label'], 'valueReference': leader['id'], 'fieldType': None, 'value': None}, 'right': {'values': expected, 'type': 'QUESTION_OPTION' if leader['type'] in CHOICE_TYPES else 'STRING', 'label': expected[0] if expected else None, 'valueReference': None, 'fieldType': None, 'value': expected[0] if expected else None}}}
            if parent:
                parent['pageItems'].append(q)
            else:
                service_map[key[0]].setdefault('questions', []).append(q)
        except ValueError as exc:
            errors.append({'sheet': 'Questions', 'row': locations[key], 'message': str(exc)})
    for key, q in questions.items():
        if q['type'] == 'RepeatableGroup' and not q['pageItems']:
            errors.append({'sheet': 'Questions', 'row': locations[key], 'message': 'Add child questions using Parent group key.'})
    import_conditions(workbook, questions, visibility, errors)
    return {name: min(number for key, number in locations.items() if key[0] == name)
            for name in {key[0] for key in locations}}


def import_conditions(workbook, questions, simple_visibility, errors):
    """Resolve a visible rule tree without storing JSON in spreadsheet cells."""
    rules = {}
    comparisons = {'Equals': 'EQUALS', 'Does not equal': 'NOT_EQUALS', 'Includes': 'INCLUDES',
                   'Does not include': 'NOT_INCLUDES', 'Starts with': 'STARTS_WITH', 'Ends with': 'ENDS_WITH'}
    for number, row in _rows(workbook, 'Conditions', errors):
        try:
            key = (row['Service name'].casefold(), row['Question key'])
            if key not in questions:
                raise ValueError('Service name and Question key must match a valid question.')
            if key in simple_visibility:
                raise ValueError('Use Conditions or the visibility columns on Questions, not both.')
            if not row['Rule key'] or row['Rule key'] in rules.get(key, {}):
                raise ValueError('Rule key must be nonempty and unique within the question.')
            parent = row['Parent rule key']
            if row['Match']:
                if row['Match'] not in ('All', 'Any') or row['Leading question key'] or row['Comparison'] or row['Expected answers']:
                    raise ValueError('Groups use Match All/Any; leave leading question, comparison, and answers blank.')
                expression = {'operator': 'AND' if row['Match'] == 'All' else 'OR',
                              'withParentheses': _boolean(row['Group parentheses'])}
            else:
                leader = questions.get((key[0], row['Leading question key']))
                if leader is None or leader is questions[key]:
                    raise ValueError('Leading question key must reference another question in this service.')
                operator = comparisons.get(row['Comparison'])
                if operator is None:
                    raise ValueError('Choose a supported Comparison.')
                values = _lines(row['Expected answers'])
                choice_label = next((option['displayValue'] for option in leader['config'].get('options', [])
                                     if values and option['dataValue'] == values[0]), values[0] if values else None)
                expression = {'operator': operator, 'withParentheses': _boolean(row['Group parentheses']),
                              'left': {'values': [], 'type': 'QUESTION_ANSWER', 'label': leader['label'],
                                       'valueReference': leader['id'], 'fieldType': None, 'value': None},
                              'right': {'values': values, 'type': 'QUESTION_OPTION' if leader['type'] in CHOICE_TYPES else 'STRING',
                                        'label': choice_label if leader['type'] in CHOICE_TYPES else None,
                                        'valueReference': None, 'fieldType': None, 'value': values[0] if values else None}}
            rules.setdefault(key, {})[row['Rule key']] = (parent, expression, number)
        except ValueError as exc:
            errors.append({'sheet': 'Conditions', 'row': number, 'message': str(exc)})
    for key, nodes in rules.items():
        number = min(v[2] for v in nodes.values())
        try:
            roots = [name for name, (parent, _, _) in nodes.items() if not parent]
            if len(roots) != 1:
                raise ValueError('Each question needs exactly one root rule (blank Parent rule key).')
            visited = set()

            def build(name, depth=0):
                if depth > 20 or name in visited:
                    raise ValueError('Condition groups contain a cycle or are too deeply nested.')
                visited.add(name)
                _, expression, _ = nodes[name]
                children = [child for child, (parent, _, _) in nodes.items() if parent == name]
                if expression['operator'] in ('AND', 'OR'):
                    if len(children) < 2:
                        raise ValueError('All/Any groups need at least two child rules.')
                    expressions = [build(child, depth + 1) for child in children]
                    result = expressions[0]
                    for child in expressions[1:]:
                        result = {**expression, 'left': result, 'right': child}
                    return result
                if children:
                    raise ValueError('Only All/Any groups can contain child rules.')
                return expression

            expression = build(roots[0])
            if len(visited) != len(nodes):
                raise ValueError('Some rules refer to a missing parent group or form a cycle.')
            questions[key]['config']['visibilityCondition'] = {'leadingQuestionExpectedAnswers': [], 'expression': expression}
        except ValueError as exc:
            errors.append({'sheet': 'Conditions', 'row': number, 'message': str(exc)})


def import_question_sheets(workbook, services, errors):
    """Expand simple inline inputs for the shared Helix importer."""
    if 'Questions' not in workbook.sheetnames:
        return _import_legacy_question_sheets(workbook, services, errors)
    headers = [c.value for c in next(workbook['Questions'].iter_rows(max_row=1, max_col=len(QUESTION_HEADERS)))]
    previous = headers[:len(PREVIOUS_QUESTION_HEADERS)] == PREVIOUS_QUESTION_HEADERS
    if headers != QUESTION_HEADERS and not previous:
        return _import_legacy_question_sheets(workbook, services, errors)
    input_headers = PREVIOUS_QUESTION_HEADERS if previous else QUESTION_HEADERS
    from openpyxl import Workbook
    expanded = Workbook()
    del expanded[expanded.active.title]
    for name, columns in LEGACY_SHEETS.items():
        expanded.create_sheet(name).append(columns)
    locations, conditions = {}, {}
    for number, cells in enumerate(workbook['Questions'].iter_rows(min_row=2, max_row=1002, max_col=101), 2):
        if all(c.value is None or str(c.value).strip() == '' for c in cells):
            continue
        try:
            if number > 1001:
                raise ValueError('Use at most 1,000 input rows per sheet.')
            if any(c.data_type in ('f', 'e') for c in cells):
                raise ValueError('Use plain values, not formulas or Excel errors.')
            if any(c.value is not None for c in cells[len(input_headers):]):
                raise ValueError('A filled cell has no matching column header.')
            row = {h: (c.value.isoformat() if isinstance(c.value, (date, datetime)) else str(c.value).strip() if c.value is not None else '') for h, c in zip(input_headers, cells)}
            if SIMPLE_TYPES.get(row['Type'], row['Type']) not in SIMPLE_TYPES.values():
                raise ValueError('Choose a supported static question type.')
            if previous:
                row['Operation'] = 'is'
                row['Value'] = row.pop('Answer is')
            if not row['Show if label'] and (row['Value'] or (not previous and row['Operation'])):
                raise ValueError('Enter Show if label for this condition.')
            if row['Show if label']:
                conditions[(row['Service name'].casefold(), row['Label'].casefold())] = (number, row['Show if label'].casefold(), row['Operation'].casefold(), row['Value'])
            legacy = {**row, 'Question key': row['Label'].casefold()}
            for col, title in enumerate(LEGACY_QUESTION_HEADERS, 1):
                expanded['Questions'].cell(number, col, legacy.get(title, ''))
            labels = _lines(row['Choices'])
            if labels and SIMPLE_TYPES.get(row['Type'], row['Type']) not in CHOICE_TYPES:
                raise ValueError('Choices are only available for choice questions.')
            for label in labels:
                expanded['Choices'].append([row['Service name'], row['Label'].casefold(), label, label, '', ''])
                locations[expanded['Choices'].max_row] = number
        except ValueError as exc:
            errors.append({'sheet': 'Questions', 'row': number, 'message': str(exc)})
    start = len(errors)
    result = _import_legacy_question_sheets(expanded, services, errors)
    for error in errors[start:]:
        if error['sheet'] == 'Choices':
            error['row'] = locations.get(error['row'], error['row'])
            error['sheet'] = 'Questions'
        error['message'] = error['message'].replace('Question key', 'Label').replace('Show when question key', 'Show if label').replace('on Choices', 'in Choices')
    questions = {(service['name'].casefold(), q['label'].casefold()): q for service in services for q in service.get('questions', [])}
    for key, (number, leader_label, operation, value) in conditions.items():
        try:
            q = questions.get(key)
            if q is None:
                continue
            leader = questions.get((key[0], leader_label))
            if leader is None or leader is q:
                raise ValueError('Show if label must reference another question in this service.')
            multi = leader['type'] in {'Checkbox', 'MultiSelectDropdown'}
            if leader['type'] not in {'TextField', 'TextArea', 'YesNo', 'Dropdown', 'RadioButtons', 'Checkbox', 'MultiSelectDropdown'}:
                raise ValueError('Conditions must reference a text or choice question.')
            operations = MULTI_OPERATIONS if multi else TEXT_OPERATIONS
            operation = operation or ('includes' if multi else 'is')
            if operation not in operations:
                raise ValueError('Supported operations for this question: ' + ', '.join(operations) + '.')
            no_value = not value or value.casefold() == '(no value)'
            options = leader['config'].get('options', [])
            selected = next((o for o in options if o['displayValue'] == value), None)
            expected = [] if no_value else [selected['dataValue'] if selected else value]
            right = {'values': expected, 'type': 'QUESTION_OPTION' if selected else 'STRING',
                     'label': selected['displayValue'] if selected else None, 'valueReference': None,
                     'fieldType': None, 'value': expected[0] if expected else None}
            q['config']['visibilityCondition'] = {'leadingQuestionId': leader['id'], 'leadingQuestionExpectedAnswers': expected,
                'expression': {'operator': operations[operation], 'withParentheses': False,
                    'left': {'values': [], 'type': 'QUESTION_ANSWER', 'label': leader['label'], 'valueReference': leader['id'], 'fieldType': None, 'value': None},
                    'right': right}}
        except ValueError as exc:
            errors.append({'sheet': 'Questions', 'row': number, 'message': str(exc)})
    expanded.close()
    return result
