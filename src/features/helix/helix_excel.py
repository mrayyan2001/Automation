"""Environment-aware XLSX templates and validated, non-persisting imports."""

import io
import zipfile

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName

from features.helix.service_generator import category_hierarchy, create_service, inspect_export
from features.helix.excel_questions import add_question_sheets, import_question_sheets

SUPPORTED = {"MULTILINE_INPUT", "RICH_INPUT", "TEXT_INPUT", "SINGLELINE_INPUT", "DROPDOWN"}


def excel_spec(data: dict) -> dict:
    info = inspect_export(data)
    fields = []
    headers = ["Service name", "Service template", "Category 1", "Category 2", "Category 3", "Tags"]
    columns = {}
    for template in info["templates"]:
        for field in template["fields"]:
            if field["type"] not in SUPPORTED or any(str(field.get(k, "")).casefold() == "platform" for k in ("name", "title")):
                continue
            header = f"{field.get('title') or field.get('name') or field['id']} [{template['id']}:{field['id']}]"
            headers.append(header)
            columns[header] = (template["id"], field["id"])
            options = ", ".join(o["name"] for o in field.get("params", {}).get("form_values", []) if isinstance(o, dict) and isinstance(o.get("name"), str))
            fields.append([template["name"], header, "Yes" if field.get("required") else "No", options])
    if len(headers) > 100:
        raise ValueError("This environment has too many profile columns for the Excel template (maximum 100).")
    categories = [[c["id"], c["path"]] for c in category_hierarchy(data["categories"])]
    return {"headers": headers, "columns": columns, "templates": info["templates"], "fields": fields,
            "categories": categories, "tags": [t["name"] for t in data.get("tags", {}).values()]}


def build_excel_template(data: dict, sample: dict | None = None) -> bytes:
    """Generate the download entirely in Python, independent of the server PATH."""
    spec = excel_spec(data)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Services"
    sheet.sheet_view.showGridLines = False
    sheet.append(spec["headers"])
    for cell in sheet[1]:
        cell.font = Font(name="Arial", size=11, bold=True, color="000000")
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    sheet.row_dimensions[1].height = 42
    for row in sheet.iter_rows(min_row=2, max_row=51, max_col=len(spec["headers"])):
        sheet.row_dimensions[row[0].row].height = 24
        for cell in row:
            cell.number_format = "@"
            cell.font = Font(name="Arial", size=11)
            cell.alignment = Alignment(vertical="center", wrap_text=True)
    for column in range(1, len(spec["headers"]) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 28 if column <= 2 else 36
    table = Table(displayName="ServicesTable", ref=f"A1:{get_column_letter(len(spec['headers']))}51")
    table.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
    sheet.add_table(table)
    sheet.freeze_panes = "A2"
    sheet.column_dimensions["B"].hidden = True

    # Keep dropdown sources in hidden columns on the same sheet, outside the service table.
    def dropdown(name, options, source_column, target):
        if not options:
            return
        for row, value in enumerate(options, 1):
            cell = sheet.cell(row, source_column, value)
            cell.data_type = "s"
            cell.number_format = "@"
        letter = get_column_letter(source_column)
        sheet.column_dimensions[letter].hidden = True
        workbook.defined_names.add(DefinedName(name, attr_text=f"'Services'!${letter}$1:${letter}${len(options)}"))
        validation = DataValidation(type="list", formula1=name, allow_blank=True)
        validation.showDropDown = False  # Excel uses False to show the list arrow.
        validation.showInputMessage = True
        validation.showErrorMessage = False  # Allow IDs or semicolon-separated categories as well.
        validation.promptTitle = "Choose an environment value"
        validation.prompt = "Select a value from the dropdown."
        sheet.add_data_validation(validation)
        validation.add(target)

    dropdown("ServiceTemplateOptions", [t["name"] for t in spec["templates"]], 102, "B2:B51")
    allowed_categories = {identifier for template in spec["templates"] for identifier in template["categoryIds"]}
    dropdown("ServiceCategoryOptions", [path for identifier, path in spec["categories"] if identifier in allowed_categories], 103, "C2:E51")
    if spec["templates"]:
        for row in range(2, 52):
            sheet.cell(row, 2, spec["templates"][0]["name"])
    add_question_sheets(workbook)
    if sample is not None:
        from features.helix.excel_sample import fill_excel_sample
        fill_excel_sample(workbook, spec, sample)
    output = io.BytesIO()
    workbook.save(output)
    workbook.close()
    return output.getvalue()


def import_excel(data: dict, contents: bytes) -> dict:
    spec = excel_spec(data)
    try:
        with zipfile.ZipFile(io.BytesIO(contents)) as archive:
            if len(archive.infolist()) > 512 or sum(member.file_size for member in archive.infolist()) > 20 * 1024 * 1024:
                raise ValueError("The Excel file is too large when unpacked.")
            if b"macroEnabled" in archive.read("[Content_Types].xml") or "xl/vbaProject.bin" in archive.namelist():
                raise ValueError("Use an .xlsx workbook without macros.")
        workbook = load_workbook(io.BytesIO(contents), read_only=True, data_only=False, keep_links=False)
    except ValueError:
        raise
    except Exception as exc:
        raise ValueError("Upload a valid .xlsx workbook using the downloaded template.") from exc
    try:
        if "Services" not in workbook.sheetnames:
            raise ValueError("The workbook needs a Services sheet.")
        sheet = workbook["Services"]
        # Only the service table area is imported. Hidden dropdown lists live beyond column 101.
        # Some XLSX writers omit dimensions, so bound iteration explicitly.
        rows = sheet.iter_rows(max_row=1001, max_col=101)
        headers = [str(c.value).strip() if c.value is not None else "" for c in next(rows, [])]
        while headers and not headers[-1]:
            headers.pop()
        if len(set(headers)) != len(headers) or any(h not in spec["headers"] and h != "Categories" for h in headers):
            raise ValueError("Headers do not match this environment. Download a new template and keep its column headers.")
        if not {"Service name", "Service template"}.issubset(headers):
            raise ValueError("Service name and Service template columns are required.")
        services, errors, seen_names = [], [], set()
        for row_number, cells in enumerate(rows, 2):
            if all(c.value is None or str(c.value).strip() == "" for index, c in enumerate(cells) if index >= len(headers) or headers[index] != "Service template"):
                continue
            if row_number > 1000 or any(c.value is not None for c in cells[100:]):
                raise ValueError("Use at most 100 columns and 1,000 worksheet rows.")
            if any(c.value is not None for c in cells[len(headers):]):
                errors.append({"sheet": "Services", "row": row_number, "message": "A filled cell has no matching column header."})
                continue
            if any(c.data_type in ("f", "e") for c in cells):
                errors.append({"sheet": "Services", "row": row_number, "message": "Use plain values, not formulas or Excel errors."})
                continue
            row = {h: str(c.value).strip() if c.value is not None else "" for h, c in zip(headers, cells)}
            try:
                template = next((t for t in spec["templates"] if row["Service template"] in (t["id"], t["name"])), None)
                if template is None:
                    raise ValueError("Choose a Service template name from the dropdown or enter its environment ID.")
                values = {}
                for header, (template_id, field_id) in spec["columns"].items():
                    if row.get(header):
                        if template_id != template["id"]:
                            raise ValueError(f"{header} belongs to a different template. Leave it blank.")
                        values[field_id] = row[header]
                category_ids = []
                category_values = [row.get(column, "") for column in ("Category 1", "Category 2", "Category 3", "Categories")]
                for value in filter(None, (c.strip() for value in category_values for c in value.split(";"))):
                    matches = [identifier for identifier, label in spec["categories"] if value in (identifier, label)]
                    if len(matches) != 1:
                        raise ValueError(f"Unknown or ambiguous category: {value}. Choose a category from the dropdown or enter its environment ID.")
                    category_ids.append(matches[0])
                category_ids = list(dict.fromkeys(category_ids))
                if len(category_ids) > 3:
                    raise ValueError("Use at most 3 categories per service.")
                tags = [t.strip() for t in row.get("Tags", "").split(",") if t.strip()]
                if len(tags) > 50:
                    raise ValueError("Use at most 50 tags per service.")
                service = {"name": row["Service name"], "templateId": template["id"], "fieldValues": values, "categoryIds": category_ids, "tags": tags}
                create_service(data, template["id"], service["name"], values, category_ids, tags)
                if service["name"].casefold() in seen_names:
                    raise ValueError("Duplicate service name in this workbook.")
                seen_names.add(service["name"].casefold())
                services.append(service)
            except ValueError as exc:
                errors.append({"sheet": "Services", "row": row_number, "message": str(exc)})
        question_rows = import_question_sheets(workbook, services, errors)
        for service in services:
            try:
                create_service(data, service["templateId"], service["name"], service["fieldValues"],
                               service["categoryIds"], service["tags"], service.get("questions", []))
            except ValueError as exc:
                errors.append({"sheet": "Questions", "row": question_rows.get(service["name"].casefold(), 1),
                               "message": f"{service['name']}: {exc}"})
        if len(services) > 50:
            raise ValueError("Import at most 50 services into one package.")
        if not services and not errors:
            raise ValueError("The Services sheet has no filled rows.")
        return {"services": [] if errors else services, "errors": errors}
    finally:
        workbook.close()
