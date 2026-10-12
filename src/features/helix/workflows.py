"""Spreadsheet workflow rows and Helix process definitions."""
from copy import deepcopy
import json
from pathlib import Path
import re
from uuid import uuid4

REFERENCE = json.loads(Path(__file__).with_name('workflow_reference.json').read_text())
LEGACY_HEADERS = ['Service name', 'Row type', 'Name', 'Type', 'Scope', 'Question label',
           'Default value', 'Target', 'Input mapping', 'Output mapping', 'From', 'To', 'Condition']
HEADERS = ['Service name', 'Approval type', 'Approval summary', 'Approvers']
APPROVALS = ['Line manager approval', 'Custom Approval']
STEPS = {'Start': 'StartEventDefinition', 'End': 'EndEventDefinition',
         'Exclusive gateway': 'ExclusiveGatewayDefinition', 'Call process': 'CallActivityDefinition',
         'Service task': 'ServiceTaskDefinition'}
SCOPES = {'Input': 'inputParams', 'Output': 'outputParams', 'Local': 'localVariables'}
NAME = re.compile(r'[A-Z][A-Za-z]*\Z')


def pascal_name(label):
    digits = ['Zero', 'One', 'Two', 'Three', 'Four', 'Five', 'Six', 'Seven', 'Eight', 'Nine']
    parts = re.findall(r'[A-Za-z]+|[0-9]+', label)
    return ''.join(''.join(digits[int(c)] for c in p) if p.isdigit()
                   else p[0].upper() + p[1:] for p in parts) or 'Question'


def check_name(name):
    if not isinstance(name, str) or not NAME.fullmatch(name):
        raise ValueError('Variable names must use PascalCase and English letters only, e.g. RequestSummary (no spaces, digits, or symbols).')


def _field(kind, name, identifier, now):
    if kind not in REFERENCE['fields']:
        raise ValueError('Variable Type must be Text, Number, Date, Object, or Context.')
    field = deepcopy(REFERENCE['fields'][kind])
    field.update(name=name, id=identifier, owner=None, lastChangedBy=None, lastUpdateTime=now)
    return field


def add_workflow_sheet(workbook):
    from openpyxl.styles import Alignment, Font
    from openpyxl.utils import get_column_letter
    from openpyxl.worksheet.datavalidation import DataValidation
    from openpyxl.worksheet.table import Table, TableStyleInfo
    sheet = workbook.create_sheet('Workflow')
    sheet.append(HEADERS)
    sheet.freeze_panes = 'B2'
    sheet.sheet_view.showGridLines = False
    for column in range(1, len(HEADERS) + 1):
        sheet.column_dimensions[get_column_letter(column)].width = 38
    for row in sheet.iter_rows(min_row=1, max_row=101, max_col=len(HEADERS)):
        sheet.row_dimensions[row[0].row].height = 42
        for cell in row:
            cell.font = Font(name='Arial', size=11, bold=cell.row == 1)
            cell.alignment = Alignment(vertical='center', wrap_text=True)
            cell.number_format = '@'
    table = Table(displayName='WorkflowTable', ref='A1:D101')
    table.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=False)
    sheet.add_table(table)
    validation = DataValidation(type='list', formula1='"' + ','.join(APPROVALS) + '"', allow_blank=True)
    validation.showDropDown = False
    validation.showErrorMessage = True
    sheet.add_data_validation(validation)
    validation.add('B2:B1001')
    guide = workbook['Guide']
    rows = [
        ['Workflow', 'One approval per row. Service name must match Services. Row order is approval order; no From/To points or variables are needed.'],
        ['Approval type', 'Choose Line manager approval or Custom Approval. Approval summary is optional. For Custom Approval, enter Approvers (the reference Flex Char 01 value).'],
        ['Approval results', 'Approvals run in row order. After each approval, Approved continues to the next approval; any other status sets the request to Rejected and ends the service. Built-in steps and status checks are added automatically.'],
    ]
    for values in rows:
        guide.append(values)
        guide.row_dimensions[guide.max_row].height = 85
        for cell in guide[guide.max_row]:
            cell.font = Font(name='Arial', size=11)
            cell.alignment = Alignment(vertical='top', wrap_text=True)


def import_workflow_sheet(workbook, services, errors):
    if 'Workflow' not in workbook.sheetnames:
        return {}
    from datetime import date, datetime
    iterator = workbook['Workflow'].iter_rows(max_row=1002, max_col=101)
    headers = [c.value for c in next(iterator)]
    columns = LEGACY_HEADERS if headers[:len(LEGACY_HEADERS)] == LEGACY_HEADERS else HEADERS
    if headers[:len(columns)] != columns or any(v is not None for v in headers[len(columns):]):
        raise ValueError('Workflow: headers do not match. Download a new template and keep its headers.')
    owners = {s['name'].casefold(): s for s in services}
    locations = {}
    for number, cells in enumerate(iterator, 2):
        if all(c.value is None or str(c.value).strip() == '' for c in cells):
            continue
        try:
            if number > 1001:
                raise ValueError('Use at most 1,000 Workflow rows.')
            if any(c.data_type in ('f', 'e') for c in cells):
                raise ValueError('Use plain values, not formulas or Excel errors.')
            if any(c.value is not None for c in cells[len(columns):]):
                raise ValueError('A filled cell has no matching column header.')
            row = {h: c.value.isoformat() if isinstance(c.value, (date, datetime)) else str(c.value).strip() if c.value is not None else '' for h, c in zip(columns, cells)}
            if columns == LEGACY_HEADERS and not row['Service name'] and row['Row type'] == 'Variable' and row['Name'] == 'Context' and row['Type'] == 'Context' and row['Scope'] == 'Input' and not any(row[h] for h in LEGACY_HEADERS[5:]):
                continue
            owner = row['Service name'].casefold()
            if owner not in owners:
                raise ValueError('Service name must match a valid row on Services.')
            if columns == LEGACY_HEADERS:
                if row['Row type'] not in ('Variable', 'Step', 'Flow'):
                    raise ValueError('Row type must be Variable, Step, or Flow.')
                if row['Row type'] == 'Variable':
                    check_name(row['Name'])
            else:
                validate_approval(row)
            owners[owner].setdefault('workflow', []).append(row)
            locations.setdefault(owner, number)
        except ValueError as exc:
            errors.append({'sheet': 'Workflow', 'row': number, 'message': str(exc)})
    return locations


def _mapping(value):
    result = []
    for line in value.splitlines():
        if not line.strip():
            continue
        target, separator, expression = line.partition('=')
        if not separator or not target.strip() or not expression.strip():
            raise ValueError('Mappings use one parameter = expression assignment per line.')
        if any(entry['assignTarget'] == target.strip() for entry in result):
            raise ValueError('Mapping targets must be unique within a step.')
        result.append({'assignTarget': target.strip(), 'expression': expression.strip()})
    return result


def _apply_advanced_workflow(workflow, rows, now):
    """Keep questionnaire input IDs while configuring variables and flow elements."""
    if not isinstance(rows, list) or len(rows) > 1000 or any(not isinstance(r, dict) for r in rows):
        raise ValueError('Workflow must contain at most 1,000 row objects.')
    rows = [{h: str(r.get(h, '')).strip() for h in LEGACY_HEADERS} for r in rows]
    variables = [r for r in rows if r['Row type'] == 'Variable']
    if any(r['Row type'] not in ('Variable', 'Step', 'Flow') for r in rows):
        raise ValueError('Workflow Row type must be Variable, Step, or Flow.')
    params = workflow['inputParams']
    question_params = {p['name'].casefold(): p for p in params}
    names = {'Context'}
    overrides = {}
    explicit_names = set()
    for row in variables:
        check_name(row['Name'])
        if row['Name'].casefold() in explicit_names:
            raise ValueError('Workflow variable names must be unique.')
        explicit_names.add(row['Name'].casefold())
        if row['Type'] == 'Context' or row['Name'] == 'Context':
            if row['Name'] != 'Context' or row['Type'] != 'Context' or row['Scope'] != 'Input' or row['Question label'] or row['Default value']:
                raise ValueError('Context is reserved as an Input variable of Type Context.')
            continue
        if row['Scope'] not in SCOPES or row['Type'] not in REFERENCE['fields']:
            raise ValueError('Choose a supported variable Scope and Type.')
        if row['Name'].casefold() == 'context':
            raise ValueError('Context is reserved.')
        names.add(row['Name'])
        if row['Question label']:
            key = row['Question label'].casefold()
            if row['Scope'] != 'Input' or key not in question_params or key in overrides:
                raise ValueError('Question label must reference a unique question input with Scope Input.')
            if question_params[key]['resourceType'] != REFERENCE['fields'][row['Type']]['resourceType']:
                raise ValueError('Variable Type must match the referenced question input type.')
            if row['Default value']:
                raise ValueError('Set question defaults on Questions, not Workflow.')
            overrides[key] = row['Name']
    used = {n.casefold() for n in names}
    for param in params:
        old = param['name'].casefold()
        name = overrides.get(old)
        if name is None:
            base = pascal_name(param['name'])
            name = base
            while name.casefold() in used:
                name += 'Next'
        else:
            used.remove(name.casefold())
        param['name'] = name
        used.add(name.casefold())
    next_id = max([536870911] + [p['id'] for p in params]) + 1
    params.append(_field('Context', 'Context', next_id, now))
    next_id += 1
    for row in variables:
        if row['Name'] == 'Context' or row['Question label']:
            continue
        field = _field(row['Type'], row['Name'], next_id, now)
        next_id += 1
        if row['Default value']:
            value = row['Default value']
            if row['Type'] == 'Number':
                try:
                    value = int(value)
                except ValueError as exc:
                    raise ValueError('Number variable defaults must be whole numbers.') from exc
            elif row['Type'] == 'Date':
                from datetime import date
                value = date.fromisoformat(value).isoformat()
            elif row['Type'] == 'Object':
                raise ValueError('Object variable defaults are not supported; use output mappings.')
            field['defaultValue'] = value
        workflow[SCOPES[row['Scope']]].append(field)
    steps = [r for r in rows if r['Row type'] == 'Step']
    flows = [r for r in rows if r['Row type'] == 'Flow']
    if not steps:
        if flows:
            raise ValueError('Add Step rows before Flow rows.')
        return
    if len(steps) > 150 or len(flows) > 500:
        raise ValueError('Use at most 150 workflow steps and 500 flows.')
    node_ids = {}
    for row in steps:
        if not row['Name'] or row['Name'] in node_ids:
            raise ValueError('Step names must be nonempty and unique.')
        if row['Type'] not in STEPS:
            raise ValueError('Choose a supported workflow step Type.')
        node_ids[row['Name']] = 'rx-' + str(uuid4())
    starts = [r['Name'] for r in steps if r['Type'] == 'Start']
    ends = [r['Name'] for r in steps if r['Type'] == 'End']
    if len(starts) != 1 or not ends:
        raise ValueError('Workflow needs exactly one Start and at least one End.')
    graph = {name: set() for name in node_ids}
    reverse = {name: set() for name in node_ids}
    for row in flows:
        source, target = row['From'], row['To']
        if source not in node_ids or target not in node_ids:
            raise ValueError('Flow From and To must match step names.')
        if source in ends or target in starts or source == target:
            raise ValueError('Flows cannot enter Start, leave End, or connect a step to itself.')
        graph[source].add(target)
        reverse[target].add(source)
    def reachable(seeds, edges):
        found, todo = set(), list(seeds)
        while todo:
            name = todo.pop()
            if name not in found:
                found.add(name)
                todo.extend(edges[name])
        return found
    if reachable(starts, graph) != set(node_ids) or reachable(ends, reverse) != set(node_ids):
        raise ValueError('All steps must be reachable from Start and able to reach an End.')
    for row in steps:
        outgoing = [f for f in flows if f['From'] == row['Name']]
        if row['Type'] == 'Exclusive gateway':
            if len(outgoing) < 2 or sum(not f['Condition'] for f in outgoing) != 1:
                raise ValueError('Each Exclusive gateway needs at least two outgoing flows and exactly one blank default Condition.')
        elif row['Type'] != 'End' and (len(outgoing) != 1 or outgoing[0]['Condition']):
            raise ValueError('Non-gateway steps need exactly one unconditional outgoing flow.')
    variable_names = {p['name'] for scope in SCOPES.values() for p in workflow[scope]}
    def expression(value):
        def activity(match):
            name = match.group(1)
            if name not in node_ids:
                raise ValueError(f'Expression references unknown step: {name}.')
            return '${activityResults.' + node_ids[name] + '.'
        value = re.sub(r'\$\{activityResults\.([^{}]+?)\.', activity, value)
        for name in re.findall(r'\$\{processContext\.([A-Za-z]+)(?:[.}])', value):
            if name not in variable_names:
                raise ValueError(f'Expression references unknown variable: {name}.')
        return value
    elements, cells = [], []
    base = {k: workflow[k] for k in ('lastUpdateTime', 'lastChangedBy', 'owner', 'tags', 'description', 'internal')}
    base['overlayGroupId'] = None
    for index, row in enumerate(steps):
        kind = STEPS[row['Type']]
        node = {**base, 'guid': node_ids[row['Name']], 'name': row['Name'],
                'resourceType': 'com.bmc.arsys.rx.services.process.domain.' + kind}
        if row['Type'] == 'End':
            node['event'] = None
        if row['Type'] in ('Call process', 'Service task'):
            if not row['Target']:
                raise ValueError('Call process and Service task steps require Target.')
            node.update(multiInstanceLoopDefinition=None, inputMap=_mapping(row['Input mapping']), outputMap=_mapping(row['Output mapping']))
            for mapping in node['inputMap'] + node['outputMap']:
                mapping['expression'] = expression(mapping['expression'])
            if any(m['assignTarget'] not in variable_names for m in node['outputMap']):
                raise ValueError('Output mapping targets must match workflow variable names.')
            if row['Type'] == 'Call process':
                node.update(calledProcessDefinitionName=row['Target'], sampleProcessDefinitionName=None)
            else:
                node.update(actionTypeName=row['Target'], runAsUser=None)
        elif row['Target'] or row['Input mapping'] or row['Output mapping']:
            raise ValueError('Only Call process and Service task steps use Target and mappings.')
        elements.append(node)
        cells.append({'id': node['guid'][3:], 'type': 'rx.' + kind.removesuffix('Definition'),
                      'position': {'x': 60 + (index % 5) * 180, 'y': 100 + (index // 5) * 150},
                      'size': {'width': 90 if row['Type'] in ('Call process', 'Service task') else 40 if row['Type'] == 'Exclusive gateway' else 30,
                               'height': 60 if row['Type'] in ('Call process', 'Service task') else 40 if row['Type'] == 'Exclusive gateway' else 30},
                      'content': row['Name'], 'z': index + 1})
    for row in flows:
        guid = 'rx-' + str(uuid4())
        elements.append({**base, 'guid': guid, 'name': 'Sequence Flow',
                         'resourceType': 'com.bmc.arsys.rx.services.process.domain.SequenceFlowDefinition',
                         'sourceNode': node_ids[row['From']], 'targetNode': node_ids[row['To']], 'condition': expression(row['Condition'])})
        cells.append({'id': guid[3:], 'type': 'rx.SequenceFlow', 'source': {'id': node_ids[row['From']][3:]},
                      'target': {'id': node_ids[row['To']][3:]}, 'z': len(cells) + 1, 'flowType': 'normal', 'labels': []})
    workflow['flowElements'] = elements
    workflow['layout'] = json.dumps({'cells': cells})


def fill_workflow_sample(workbook, service_name):
    sheet = workbook['Workflow']
    for number, values in enumerate([
        [service_name, 'Line manager approval', 'Line Manager Approval', ''],
        [service_name, 'Custom Approval', '', 'Max'],
    ], 2):
        for column, value in enumerate(values, 1):
            cell = sheet.cell(number, column, value)
            cell.data_type = 's'


def validate_approval(row):
    if row.get('Approval type') == 'Ad hoc approval':
        row['Approval type'] = 'Custom Approval'
    if row.get('Approval type') not in APPROVALS:
        raise ValueError('Choose Line manager approval or Custom Approval.')
    if row['Approval type'] == 'Custom Approval' and not row.get('Approvers', '').strip():
        raise ValueError('Custom Approval requires Approvers.')
    if row['Approval type'] == 'Line manager approval' and row.get('Approvers', '').strip():
        raise ValueError('Line manager approval determines the manager automatically; leave Approvers blank.')


def apply_workflow(workflow, rows, now):
    """Build the standard service pipeline from an ordered approval list."""
    if not isinstance(rows, list) or len(rows) > 1000 or any(not isinstance(r, dict) for r in rows):
        raise ValueError('Workflow must contain at most 1,000 row objects.')
    if rows and any('Row type' in row for row in rows):
        _apply_advanced_workflow(workflow, rows, now)
        _ensure_qanswer(workflow, now)
        return
    rows = [{h: '' if r.get(h) is None else str(r.get(h, '')).strip() for h in HEADERS} for r in rows]
    if len(rows) > 30:
        raise ValueError('Use at most 30 approval steps per service.')
    expanded = [
        {'Row type': 'Step', 'Name': 'Start', 'Type': 'Start'},
        {'Row type': 'Step', 'Name': 'QAnswer', 'Type': 'Call process',
         'Target': next(iter(REFERENCE['dependencies'].values()))['name'],
         'Input mapping': 'SBContext = ${processContext.Context}'},
        {'Row type': 'Flow', 'From': 'Start', 'To': 'QAnswer'},
    ]
    previous = 'QAnswer'
    for index, row in enumerate(rows):
        validate_approval(row)
        kind = row['Approval type']
        reference = next(e for e in REFERENCE['workflow']['flowElements']
                         if e.get('name') == ('manager approval' if kind == 'Line manager approval' else 'Ad-Hoc Approval'))
        name = ('LineManagerApproval' if kind == 'Line manager approval' else 'CustomApproval') + 'Next' * index
        mappings = deepcopy(reference['inputMap'])
        for mapping in mappings:
            if mapping['assignTarget'] == 'Approval Summary' and row.get('Approval summary', '').strip():
                mapping['expression'] = json.dumps(row['Approval summary'].strip())
            if mapping['assignTarget'] == 'Flex Char 01':
                mapping['expression'] = json.dumps(row['Approvers'].strip())
        check = name + 'Status'
        rejected = name + 'Rejected'
        expanded.extend([
            {'Row type': 'Step', 'Name': name, 'Type': 'Call process', 'Target': reference['calledProcessDefinitionName'],
             'Input mapping': '\n'.join(m['assignTarget'] + ' = ' + m['expression'] for m in mappings)},
            {'Row type': 'Flow', 'From': previous, 'To': name},
            {'Row type': 'Step', 'Name': check, 'Type': 'Exclusive gateway'},
            {'Row type': 'Flow', 'From': name, 'To': check},
            {'Row type': 'Step', 'Name': rejected, 'Type': 'Service task', 'Target': 'myit-sb:setServiceRequestStatus',
             'Input mapping': 'requestId = ${processContext.Context.serviceRequest.id}\nstatus = "Rejected"'},
            {'Row type': 'Flow', 'From': check, 'To': rejected},
            {'Row type': 'Step', 'Name': rejected + 'End', 'Type': 'End'},
            {'Row type': 'Flow', 'From': rejected, 'To': rejected + 'End'},
        ])
        # The approved branch joins the following step; its condition refers to this approval.
        previous = check
    expanded.append({'Row type': 'Step', 'Name': 'End', 'Type': 'End'})
    expanded.append({'Row type': 'Flow', 'From': previous, 'To': 'End'})
    for row in expanded:
        if row['Row type'] == 'Flow' and row['From'].endswith('Status') and not row['To'].endswith('Rejected'):
            approval = row['From'].removesuffix('Status')
            row['Condition'] = '${activityResults.' + approval + '.Status} = "Approved"'
    _apply_advanced_workflow(workflow, expanded, now)
    # Keep the approved path on the top row and each rejection branch below its gateway.
    positions = {'Start': (60, 100), 'QAnswer': (240, 100), 'End': (420 + len(rows) * 360, 100)}
    for index, row in enumerate(rows):
        name = ('LineManagerApproval' if row['Approval type'] == 'Line manager approval' else 'CustomApproval') + 'Next' * index
        positions.update({name: (420 + index * 360, 100), name + 'Status': (600 + index * 360, 110),
                          name + 'Rejected': (575 + index * 360, 260), name + 'RejectedEnd': (605 + index * 360, 410)})
    nodes = {e['guid'][3:]: e for e in workflow['flowElements']}
    layout = json.loads(workflow['layout'])
    for cell in layout['cells']:
        element = nodes[cell['id']]
        if element['name'] in positions:
            x, y = positions[element['name']]
            cell['position'] = {'x': x, 'y': y}
            if element['resourceType'].endswith('EndEventDefinition'):
                label = 'End (rejected)' if element['name'].endswith('RejectedEnd') else 'End (approved)' if rows else 'End'
                element['name'] = label
                cell['content'] = label
        elif element.get('sourceNode'):
            source = nodes[element['sourceNode'][3:]]
            if source['resourceType'].endswith('ExclusiveGatewayDefinition'):
                cell['labels'] = [{'attrs': {'text': {'text': 'Approved' if element['condition'] else 'Not approved'}}, 'position': 0.5}]
    workflow['layout'] = json.dumps(layout)


def _ensure_qanswer(workflow, now):
    """Normalize older workflow graphs to the required Start → QAnswer entry."""
    elements = workflow['flowElements']
    starts = [e for e in elements if e['resourceType'].endswith('StartEventDefinition')]
    start = starts[0]
    target = next(iter(REFERENCE['dependencies'].values()))['name']
    calls = [e for e in elements if e.get('calledProcessDefinitionName') == target]
    if len(calls) > 1:
        raise ValueError('Workflow must have only one QAnswer step.')
    if calls:
        call = calls[0]
        outgoing = [e for e in elements if e.get('sourceNode') == call['guid']]
        if len(outgoing) != 1 or outgoing[0].get('condition'):
            raise ValueError('QAnswer requires one unconditional outgoing flow.')
        successor = outgoing[0]['targetNode']
        for e in elements:
            if e.get('targetNode') == call['guid']:
                e['targetNode'] = successor
        elements.remove(outgoing[0])
    else:
        reference = next(e for e in REFERENCE['workflow']['flowElements'] if e.get('name') == 'QAnswers')
        call = deepcopy(reference)
        call.update(guid='rx-' + str(uuid4()), owner=None, lastChangedBy=None, lastUpdateTime=now)
        elements.append(call)
    call['name'] = 'QAnswer'
    call['inputMap'] = [{'assignTarget': 'SBContext', 'expression': '${processContext.Context}'}]
    initial = next(e for e in elements if e.get('sourceNode') == start['guid'])
    successor = initial['targetNode']
    initial['targetNode'] = call['guid']
    continuation = {**initial, 'guid': 'rx-' + str(uuid4()), 'sourceNode': call['guid'], 'targetNode': successor}
    elements.append(continuation)
    layout = json.loads(workflow['layout'])
    layout['cells'] = [c for c in layout['cells'] if c['id'] in {e['guid'][3:] for e in elements}]
    for cell in layout['cells']:
        element = next(e for e in elements if e['guid'][3:] == cell['id'])
        if element.get('sourceNode'):
            cell['source'] = {'id': element['sourceNode'][3:]}
            cell['target'] = {'id': element['targetNode'][3:]}
    if not any(c['id'] == call['guid'][3:] for c in layout['cells']):
        layout['cells'].append({'id': call['guid'][3:], 'type': 'rx.CallActivity', 'position': {'x': 180, 'y': 100},
                                'size': {'width': 90, 'height': 60}, 'content': 'QAnswer', 'z': len(layout['cells']) + 1})
    layout['cells'].append({'id': continuation['guid'][3:], 'type': 'rx.SequenceFlow', 'source': {'id': call['guid'][3:]},
                            'target': {'id': successor[3:]}, 'z': len(layout['cells']) + 1, 'flowType': 'normal', 'labels': []})
    workflow['layout'] = json.dumps(layout)


def workflow_dependencies(workflow, environment, now):
    """Include referenced custom processes instead of leaving dangling call activities."""
    available = {p['name']: p for p in REFERENCE['dependencies'].values()}
    available.update({p['name']: p for p in environment.get('processDefinitionGuidToProcessDefinition', {}).values()})
    result = {}
    def include(elements):
        for element in elements:
            target = element.get('calledProcessDefinitionName')
            if target in available and target not in result:
                process = deepcopy(available[target])
                result[target] = process
                include(process['flowElements'])
            if element.get('flowElements'):
                include(element['flowElements'])
    include(workflow['flowElements'])
    process_names = {name: name + ' ' + workflow['guid'][3:11] for name in result}
    def update_calls(elements):
        for element in elements:
            target = element.get('calledProcessDefinitionName')
            if target in process_names:
                element['calledProcessDefinitionName'] = process_names[target]
            update_calls(element.get('flowElements', []))
    update_calls(workflow['flowElements'])
    interfaces = {}
    for name, process in result.items():
        process['name'] = process_names[name]
        update_calls(process['flowElements'])
        used, replacements = set(), {}
        for scope in SCOPES.values():
            for field in process.setdefault(scope, []):
                original = field['name']
                new = pascal_name(original)
                while new.casefold() in used:
                    new += 'Next'
                field['name'] = new
                used.add(new.casefold())
                replacements[original] = new
        interfaces[process['name']] = {old: new for old, new in replacements.items() if any(f['name'] == new for f in process['inputParams'])}
        if 'context' not in used:
            identifier = max([536870911] + [f['id'] for scope in SCOPES.values() for f in process[scope]]) + 1
            process['inputParams'].append(_field('Context', 'Context', identifier, now))
        def rename(value):
            if isinstance(value, dict):
                return {k: replacements.get(v, v) if k == 'assignTarget' and isinstance(v, str) else rename(v) for k, v in value.items()}
            if isinstance(value, list):
                return [rename(v) for v in value]
            if isinstance(value, str):
                for old, new in replacements.items():
                    value = value.replace('processContext.' + old + '}', 'processContext.' + new + '}').replace('processContext.' + old + '.', 'processContext.' + new + '.')
            return value
        # Output assignment targets are workflow variables; input targets are callee parameters.
        def rename_elements(elements):
            for element in elements:
                for key in ('condition', 'inputMap', 'outputMap'):
                    if key in element:
                        if key == 'inputMap':
                            for mapping in element[key]:
                                mapping['expression'] = rename(mapping['expression'])
                        else:
                            element[key] = rename(element[key])
                rename_elements(element.get('flowElements', []))
        rename_elements(process['flowElements'])
    def update_inputs(elements):
        for element in elements:
            mapping = interfaces.get(element.get('calledProcessDefinitionName'), {})
            for entry in element.get('inputMap', []):
                entry['assignTarget'] = mapping.get(entry['assignTarget'], entry['assignTarget'])
            update_inputs(element.get('flowElements', []))
    update_inputs(workflow['flowElements'])
    for process in result.values():
        update_inputs(process['flowElements'])
    # Fresh process and node identities; replacements cover expressions and layout too.
    ids = {}
    for process in result.values():
        def collect(value):
            if isinstance(value, dict):
                if isinstance(value.get('guid'), str):
                    ids.setdefault(value['guid'], 'rx-' + str(uuid4()))
                for child in value.values():
                    collect(child)
            elif isinstance(value, list):
                for child in value:
                    collect(child)
        collect(process)
    def clone(value):
        if isinstance(value, dict):
            return {k: None if k in ('owner', 'lastChangedBy') else now if k == 'lastUpdateTime' else clone(v) for k, v in value.items()}
        if isinstance(value, list):
            return [clone(v) for v in value]
        if isinstance(value, str):
            for old, new in ids.items():
                value = value.replace(old, new).replace(old.removeprefix('rx-'), new.removeprefix('rx-'))
        return value
    return [clone(p) for p in result.values()]
