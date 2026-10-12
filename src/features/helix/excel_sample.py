"""Flatten the static service sample into editable workbook input rows."""
from copy import deepcopy
from html.parser import HTMLParser
from uuid import uuid4

from openpyxl.styles import Alignment, Font

from features.helix.excel_questions import SIMPLE_TYPES, MULTI_TYPES, SHEETS, TEXT_OPERATIONS, MULTI_OPERATIONS


class _PlainText(HTMLParser):
    def __init__(self):
        super().__init__()
        self.parts = []

    def handle_data(self, data):
        self.parts.append(data)

    def handle_starttag(self, tag, attrs):
        if tag == 'br':
            self.parts.append('\n')

    def handle_endtag(self, tag):
        if tag in ('p', 'div'):
            self.parts.append('\n')


def _text(value):
    parser = _PlainText()
    parser.feed(value or '')
    return ''.join(parser.parts).strip()


def fill_excel_sample(workbook, spec, sample):
    sample = deepcopy(sample)
    # The reference contains a dynamic-only single dropdown. Use its static
    # checkbox choices for a connector-free example of that supported type.
    choice = next(q for q in sample['questions'] if q['type'] == 'Checkbox')
    sample['questions'].append({'id': str(uuid4()), 'type': 'Dropdown',
                               'label': 'Static single choice dropdown', 'required': False,
                               'config': {'options': deepcopy(choice['config']['options']),
                                          'defaultAnswer': choice['config']['options'][0]['dataValue']}})
    template = next(t for t in spec['templates'] if t['id'] == sample['templateId'])
    category_paths = dict(spec['categories'])

    def write(name, number, values):
        sheet = workbook[name]
        headers = spec['headers'] if name == 'Services' else SHEETS[name]
        for column, header in enumerate(headers, 1):
            if header not in values or values[header] is None:
                continue
            cell = sheet.cell(number, column, values[header])
            if isinstance(values[header], str):
                cell.data_type = 's'

    service_row = {'Service name': sample['name'], 'Service template': template['name'], 'Tags': ', '.join(sample.get('tags', []))}
    for index, identifier in enumerate(sample.get('categoryIds', []), 1):
        service_row[f'Category {index}'] = category_paths[identifier]
    for header, (template_id, field_id) in spec['columns'].items():
        if template_id == template['id']:
            service_row[header] = sample['fieldValues'].get(field_id, '')
    write('Services', 2, service_row)
    workbook['Services'].row_dimensions[2].height = 45

    questions = [q for q in sample['questions'] if q['type'] in SIMPLE_TYPES.values()]
    used_labels = set()
    for q in questions:
        original = q['label']
        suffix = 2
        while q['label'].casefold() in used_labels:
            q['label'] = f'{original} ({suffix})'
            suffix += 1
        used_labels.add(q['label'].casefold())
    labels = {q['id']: q['label'] for q in questions}
    type_names = {value: name for name, value in SIMPLE_TYPES.items()}

    def limit(value):
        if not value or value.get('limitType') == 'NO_LIMIT':
            return ''
        if value.get('limitType') == 'SPECIFIC':
            return value['specificLimit']
        return f"today{value['relativeLimit']['numberOfUnits']:+d}"

    for number, q in enumerate(questions, 2):
        c = q.get('config', {})
        row = {'Service name': sample['name'], 'Label': q['label'], 'Type': type_names[q['type']],
               'Required': 'Yes' if q.get('required') else 'No'}
        for header, setting in [('Hidden', 'hidden'), ('Read-only', 'readOnly')]:
            if setting in c:
                row[header] = 'Yes' if c[setting] else 'No'
        options = c.get('options', [])
        row['Choices'] = '\n'.join(o['displayValue'] for o in options)
        values = {o['dataValue']: o['displayValue'] for o in options}
        defaults = c.get('defaultAnswers', []) if q['type'] in MULTI_TYPES else [c['defaultAnswer']] if c.get('defaultAnswer') is not None else []
        row['Default answer'] = '\n'.join(str(values.get(v, v)) for v in defaults)
        row['Earliest date'] = limit(c.get('min'))
        row['Latest date'] = limit(c.get('max'))
        expression = (c.get('visibilityCondition') or {}).get('expression', {})
        operations = {value: name for name, value in {**TEXT_OPERATIONS, **MULTI_OPERATIONS}.items()}
        if expression.get('operator') in operations:
            leader_id = expression['left']['valueReference']
            if leader_id in labels:
                leader = next(item for item in questions if item['id'] == leader_id)
                mapping = {o['dataValue']: o['displayValue'] for o in leader.get('config', {}).get('options', [])}
                expected = expression['right'].get('values') or [expression['right'].get('value')]
                row['Show if label'] = labels[leader_id]
                row['Operation'] = operations[expression['operator']]
                row['Value'] = next((str(mapping.get(v, v)) for v in expected if v is not None), '(no value)')
        write('Questions', number, row)
        workbook['Questions'].row_dimensions[number].height = max(36, 16 * len(options) + 8)
    from features.helix.workflows import fill_workflow_sample
    fill_workflow_sample(workbook, sample['name'])
    guide = workbook['Guide']
    guide.append(['Filled sample', 'Supported question types are prefilled. Choices use their labels as saved values. Single visibility conditions are included; tables, repeatable groups, and compound rules are omitted.'])
    guide.row_dimensions[guide.max_row].height = 70
    for cell in guide[guide.max_row]:
        cell.font = Font(name='Arial', size=11)
        cell.alignment = Alignment(vertical='top', wrap_text=True)
